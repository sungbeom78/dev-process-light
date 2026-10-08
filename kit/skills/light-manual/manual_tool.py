#!/usr/bin/env python3
"""사용설명서 도우미 (AI 가 실행한다. 사용자가 직접 쓸 필요는 없음). 파이썬만 있으면 check·outline 은 바로 된다.

    python manual_tool.py check   manual/        # 규칙 점검: 페이지마다 그림·대본, 한 단계 = 한 페이지, 쪽수, 글자 수, 링크
    python manual_tool.py review  manual/        # 모든 페이지를 그려 넘침 찾기 + 장별 미리보기 PNG (.manual-review/)
    python manual_tool.py outline manual/ [장]   # 채팅에 보여 줄 개요 (페이지별 제목·핵심)
    python manual_tool.py pptx    manual/ out.pptx   # pptx: 페이지 = 슬라이드 한 장, 대본 = 발표자 노트
                                                     #  (pip install playwright python-pptx && python -m playwright install chromium)
"""
from __future__ import annotations
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

import hashlib
import json

VISUAL_CLASSES = {"flow", "cards", "compare", "screen", "chat", "tree", "diagram", "big",
                  "term", "lanes", "stack", "timeline", "kv", "check"}
VISUAL_TAGS = {"svg", "img", "table", "video"}
NEED = {"index": ["cover", "start-step", "chapters", "changes"], "feature": ["overview", "step"]}
MIN_PAGES = {"index": 6, "feature": 5}          # 바닥선 (목표 아님). 기능 장은 표지 + 흐름 + 단계마다 한 쪽
SHOW = {"screen", "img", "term"}                # 단계 페이지에 있어야 하는 것: 화면 흉내·캡처·터미널
MAX_CHARS, MIN_NOTES = 350, 60
REVIEW = ".review.json"
VOID = {"meta", "link", "br", "img", "input", "hr", "source", "path", "rect", "circle", "line", "polyline", "polygon", "ellipse", "use"}


class Deck(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title, self.slides, self.links, self.assets = "", [], [], []
        self.stack, self.cur, self.in_title, self.head = [], None, False, 0

    def flags(self):
        return self.stack[-1][1:] if self.stack else (False, False)

    def handle_starttag(self, tag, attrs):
        a = {k: (v or "") for k, v in attrs}
        notes, hidden = self.flags()
        cls = set(a.get("class", "").split())
        self.in_title = self.in_title or tag == "title"
        if tag == "section" and "slide" in cls:
            self.cur = {"title": a.get("data-title", ""), "kind": a.get("data-kind", ""), "text": [], "notes": [], "head": [],
                        "visual": False, "feat": set(), "flows": []}
            self.slides.append(self.cur)
        if self.cur is not None and (tag in ("h1", "h2") or "lead" in cls) and not self.head:
            self.head = len(self.stack) + 1
            self.cur["head"].append("\x00")
        if self.cur is not None:
            notes = notes or (tag == "aside" and "notes" in cls)
            hidden = hidden or tag in ("script", "style")
            if not notes and (tag in VISUAL_TAGS or cls & VISUAL_CLASSES):
                self.cur["visual"] = True
            if not notes:
                self.cur["feat"].add(tag); self.cur["feat"].update(cls)
                top = self.stack[-1] if self.stack else None
                if tag in ("ol", "ul") and "flow" in cls:
                    self.cur["flows"].append({"depth": len(self.stack) + 1, "items": []})
                elif tag == "li" and self.cur["flows"] and self.cur["flows"][-1]["depth"] == len(self.stack):
                    self.cur["flows"][-1]["items"].append(False)
                elif tag in ("span", "small", "code") and self.cur["flows"] and self.cur["flows"][-1]["items"] \
                        and len(self.stack) > self.cur["flows"][-1]["depth"]:
                    self.cur["flows"][-1]["items"][-1] = True
                del top
        for k in ("href", "src"):
            if a.get(k):
                self.links.append(a[k])
        if (tag == "link" and "stylesheet" in a.get("rel", "")) or (tag == "script" and a.get("src")):
            self.assets.append(a.get("href") or a.get("src"))
        if tag not in VOID:
            self.stack.append((tag, notes, hidden))

    def handle_endtag(self, tag):
        self.in_title = False if tag == "title" else self.in_title
        while self.stack:
            if self.stack.pop()[0] == tag:
                break
        if self.head and len(self.stack) < self.head:
            self.head = 0
        if tag == "section" and not any(t[0] == "section" for t in self.stack):
            self.cur = None

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        if self.cur is not None and not self.flags()[1]:
            (self.cur["notes"] if self.flags()[0] else self.cur["text"]).append(data)
            if self.head and not self.flags()[0]:
                self.cur["head"].append(data)


def parse(f: Path) -> Deck:
    d = Deck()
    d.feed(f.read_text(encoding="utf-8"))
    for s in d.slides:
        s["text"] = re.sub(r"\s+", " ", " ".join(s["text"])).strip()
        s["notes"] = re.sub(r"\s+", " ", " ".join(s["notes"])).strip()
        s["head"] = " / ".join(x for x in (re.sub(r"\s+", " ", b).strip() for b in "".join(s["head"]).split("\x00")) if x)
    return d


def check(mdir: Path) -> list[str]:
    probs = []
    for a in ("manual.css", "manual.js"):
        if not (mdir / "assets" / a).is_file():
            probs.append(f"assets/{a} 없음 -- light-manual 스킬의 assets 를 복사")
    if not (mdir / "index.html").is_file():
        probs.append("index.html(목차) 없음")
    linked = set()
    for f in sorted(mdir.glob("*.html")):
        d = parse(f)
        raw = f.read_text(encoding="utf-8")
        if re.search(r"<!--\s*작성\s*:|\{\{[a-z_]+\}\}", raw):
            probs.append(f"{f.name}: 채우지 않은 자리표시가 남아 있음")
        for a in d.assets:
            if re.match(r"^(https?:)?//", a):
                probs.append(f"{f.name}: 인터넷 주소 사용 금지 ({a})")
        for u in d.links:
            if re.match(r"^[a-z][a-z0-9+.-]*:", u, re.I) or u.startswith(("#", "//")):
                continue
            if f.name == "index.html" and u.split("#")[0].endswith(".html"):
                linked.add(Path(u.split("#")[0]).name)
            if not (f.parent / u.split("#")[0].split("?")[0]).exists():
                probs.append(f"{f.name}: 깨진 링크 {u}")
        role = "index" if f.name == "index.html" else "feature"
        for i, s in enumerate(d.slides, 1):
            w = f"{f.name} #{i} '{s['title']}'"
            if not s["title"]:
                probs.append(f"{w}: data-title 없음")
            if len(s["notes"]) < MIN_NOTES:
                probs.append(f"{w}: 대본(aside.notes) 없음/짧음")
            if s["kind"] not in ("cover", "section") and not s["visual"]:
                probs.append(f"{w}: 그림 없음 (순서도·카드·비교·화면·표·svg)")
            if len(s["text"].replace(" ", "")) > MAX_CHARS:
                probs.append(f"{w}: 글이 너무 많음 -- 나누거나 그림으로")
            if s["kind"] in ("step", "start-step") and not (s["feat"] & SHOW):
                probs.append(f"{w}: 따라 하기 페이지에는 화면(.screen)·캡처(img)·터미널(.term) 중 하나")
            for fl in s["flows"]:
                bare = sum(1 for x in fl["items"] if not x)
                if bare:
                    probs.append(f"{w}: 순서도 칸 {bare}개에 설명(<span>) 없음 -- 제목만 있는 상자 금지")
        kinds = {s["kind"] for s in d.slides}
        probs += [f"{f.name}: 필요한 페이지 없음 data-kind=\"{k}\"" for k in NEED[role] if k not in kinds]
        need = max(MIN_PAGES[role], 2 + sum(1 for s in d.slides if s["kind"] == "step")) if role == "feature" else MIN_PAGES[role]
        if len(d.slides) < need:
            probs.append(f"{f.name}: {len(d.slides)}쪽 -- 최소 {need}쪽 (표지·흐름·단계마다 한 쪽)")
    for f in sorted(mdir.glob("*.html")):
        if f.name != "index.html" and f.name not in linked:
            probs.append(f"{f.name}: 목차(index.html)에서 연결되지 않음")
    rf = mdir / REVIEW
    if rf.is_file():                                   # 렌더링 검토를 했다면: 그 뒤 바뀌지 않았고 넘침이 없어야 한다
        try:
            rec = json.loads(rf.read_text(encoding="utf-8"))
        except ValueError:
            rec = {}
        if rec.get("files") != fingerprints(mdir):
            probs.append("렌더링 검토 뒤 매뉴얼이 바뀜 -- manual_tool.py review 다시 (미리보기를 다시 눈으로)")
        probs += [f"렌더링: {o}" for o in rec.get("overflow", [])]
    return probs


def fingerprints(mdir: Path) -> dict:
    files = sorted(list(mdir.glob("*.html")) + list((mdir / "assets").glob("*")))
    return {f.relative_to(mdir).as_posix(): hashlib.sha256(f.read_bytes().replace(b"\r\n", b"\n")).hexdigest()[:16] for f in files if f.is_file()}


OVERFLOW_JS = """() => Array.from(document.querySelectorAll('section.slide')).map((sl, i) => {
  const s = sl.querySelector(':scope > .s') || sl, sr = sl.getBoundingClientRect();
  const bodies = Array.from(s.querySelectorAll('.body')), all = Array.from(s.querySelectorAll('*'));
  const over = Math.max(s.scrollHeight - s.clientHeight, ...bodies.map(b => b.scrollHeight - b.clientHeight));
  const head = s.querySelector(':scope > h1, :scope > h2'), hb = head ? head.getBoundingClientRect().bottom : -1e9;
  return {i: i + 1, title: sl.dataset.title || '', over: over,
          wide: all.some(e => e.getBoundingClientRect().right > sr.right + 2),
          low: all.some(e => e.getBoundingClientRect().bottom > sr.bottom + 2),
          up: bodies.some(b => Array.from(b.children).some(e => e.getBoundingClientRect().top < hb - 2))};
})"""


def review(mdir: Path) -> dict:
    """모든 페이지를 브라우저로 그려 넘침을 찾고 장별 한 장짜리 미리보기(PNG)를 만든다. 결과는 manual/.review.json."""
    from playwright.sync_api import sync_playwright
    out = mdir.resolve().parent / ".manual-review"
    out.mkdir(exist_ok=True)
    (out / ".gitignore").write_text("*\n", encoding="utf-8")      # 미리보기는 저장(커밋)하지 않는다
    overflow, sheets, pages = [], [], 0
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1320, "height": 800})
        for f in sorted(mdir.glob("*.html"), key=lambda x: (x.name != "index.html", x.name)):
            pg.set_viewport_size({"width": 1320, "height": 800})
            pg.goto(f.resolve().as_uri() + "?export"); pg.wait_for_timeout(250)
            res = pg.evaluate(OVERFLOW_JS)
            pages += len(res)
            for r in res:
                why = (["아래로 넘침 %dpx" % r["over"]] if r["over"] > 4 else ["페이지 아래로 삐져나감"] if r["low"] else []) + \
                      (["옆으로 넘침"] if r["wide"] else []) + (["본문이 제목을 덮음"] if r["up"] else [])
                if why:
                    overflow.append(f"{f.name} #{r['i']} '{r['title']}': " + ", ".join(why) + " -- 줄이거나 페이지를 나눈다")
            pg.set_viewport_size({"width": 1930, "height": 1000})
            pg.goto(f.resolve().as_uri() + "?sheet"); pg.wait_for_timeout(250)
            shot = out / f"{f.stem}.png"; pg.screenshot(path=str(shot), full_page=True)
            sheets.append(shot.relative_to(mdir.resolve().parent).as_posix())
        b.close()
    rec = {"pages": pages, "files": fingerprints(mdir), "overflow": overflow, "sheets": sheets}
    (mdir / REVIEW).write_text(json.dumps(rec, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return rec


def outline(mdir: Path, which: str | None) -> str:
    files = sorted(mdir.glob("*.html"), key=lambda p: (p.name != "index.html", p.name))
    if which:
        files = [f for f in files if which.lower() in (f.stem.lower(), parse(f).title.lower()) or which in parse(f).title]
    out = []
    for f in files:
        d = parse(f)
        out.append(f"### {d.title.split(' — ')[0]}  ·  `{f.as_posix()}`  ({len(d.slides)}쪽)")
        for i, s in enumerate(d.slides, 1):
            t = s["head"] or s["text"]
            out.append(f"{i}. **{s['title']}** — {t[:160]}")
        out.append("")
    out.append(f"열기: `{(mdir / 'index.html').as_posix()}` (브라우저 · P 발표 · N 대본 · 인쇄 = 16:9 PDF)")
    return "\n".join(out)


def pptx(mdir: Path, out: Path) -> int:
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    tmp = out.with_suffix(".pages"); tmp.mkdir(parents=True, exist_ok=True)
    n = 0
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={"width": 1320, "height": 800}, device_scale_factor=1.5)
        for f in sorted(mdir.glob("*.html"), key=lambda p: (p.name != "index.html", p.name)):
            d = parse(f)
            pg.goto(f.resolve().as_uri() + "?export"); pg.wait_for_timeout(200)
            for i, el in enumerate(pg.query_selector_all("section.slide")):
                n += 1; img = tmp / f"{n:03d}.png"; el.screenshot(path=str(img))
                sl = prs.slides.add_slide(prs.slide_layouts[6])
                sl.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
                s = d.slides[i] if i < len(d.slides) else {"title": "", "notes": ""}
                sl.notes_slide.notes_text_frame.text = f"{s['title']}\n{s['notes']}".strip()
        b.close()
    prs.save(str(out))
    for x in tmp.iterdir():
        x.unlink()
    tmp.rmdir()
    return n


def main(argv: list[str]) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass
    if len(argv) < 2 or argv[0] not in ("check", "outline", "pptx", "review"):
        print(__doc__); return 2
    mdir = Path(argv[1])
    if argv[0] == "check":
        probs = check(mdir)
        print("\n".join(f"✗ {p}" for p in probs) or "✓ 사용설명서 규칙 통과")
        return 1 if probs else 0
    if argv[0] == "review":
        try:
            rec = review(mdir)
        except ImportError:
            print("미리보기에는 playwright 가 필요합니다: pip install playwright && python -m playwright install chromium\n"
                  "없으면 manual/index.html 을 브라우저로 열어 페이지마다 글이 넘치지 않는지 눈으로 봅니다."); return 77
        print("\n".join(f"✗ {o}" for o in rec["overflow"]) or f"✓ {rec['pages']}쪽 넘침 없음")
        print("미리보기: " + " · ".join(rec["sheets"]) + "  -- 열어서 눈으로 확인하고 채팅에도 보여 준다")
        return 1 if rec["overflow"] else 0
    if argv[0] == "outline":
        print(outline(mdir, argv[2] if len(argv) > 2 else None)); return 0
    out = Path(argv[2] if len(argv) > 2 else "manual.pptx")
    print(f"pptx: {out} ({pptx(mdir, out)}쪽)"); return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
