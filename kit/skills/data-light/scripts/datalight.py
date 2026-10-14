#!/usr/bin/env python3
"""data-light -- easy data design for small systems (data-expert-framework-light).

One JSON file (data/model.light.json) describes your tables. From it this tool
checks the design, builds SQL / ERD / spec, tries it in SQLite, imports csv/xlsx,
and upgrades to the full data-expert-framework model.json.

Commands
  check   <model>                         check the design (exit 0 ok / 1 problems)
  split   <model> [--write]               move personal-info columns into pi_<table>
  build   <model> [--out DIR] [--db sqlite|postgresql]
                                          schema.sql, erd.md, spec.md (+ spec.xlsx if openpyxl)
  try     <model>                         run the schema + sample rows in SQLite and probe the rules
  import plan  <file> --table NAME [--label 한글] [--sheet S] [--out plan.json]
  import apply <plan.json> --model <model> --db data/app.db --approval "<사용자 답 그대로>"
  upgrade <model> [--out model.json] [--schema app] [--pi-schema pi]
                                          convert to the full framework's model.json
  verify  <model> [--db app.db|--no-db] [--mutate]
                                          one pass/fail gate: check + goals + rules + out/ fresh + DB safety
  hook install [--repo .] [--no-ci] [--append]
                                          run `verify` on every commit that changes a model (+ CI workflow)
  hook status [--repo .]                  installed hook / tool version vs this one (exit 0 same, 1 differ)
  import schema <app.db> [--out data/model.light.json] [--force]
                                          draft a model from an existing SQLite DB (labels left empty)
  --version

Harness idea: decide "what counts as done" first (goals / not_doing in the model), turn it into
runnable checks (try / verify), and run them automatically at commit time (hook) and on push (CI).

Python standard library only (openpyxl is optional: xlsx import / spec.xlsx).
MIT License -- see LICENSE.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import os
import re
import sqlite3
import sys
from pathlib import Path

FORMAT = "data-light/1"
__version__ = "0.4.0"
NAME = re.compile(r"^[a-z][a-z0-9_]{0,62}$")
TYPES = ("id", "ref", "text", "long_text", "int", "number", "money", "bool", "date", "datetime", "owner")
ACCESS = ("members", "own", "read")
RESERVED = {
    "all", "and", "any", "as", "asc", "between", "by", "case", "check", "column", "constraint", "create", "cross",
    "current_date", "current_time", "current_timestamp", "default", "delete", "desc", "distinct", "drop", "else", "end",
    "except", "exists", "false", "for", "foreign", "from", "full", "grant", "group", "having", "in", "index", "inner",
    "insert", "intersect", "into", "is", "join", "key", "left", "like", "limit", "natural", "not", "null", "offset", "on",
    "or", "order", "outer", "primary", "references", "right", "select", "set", "table", "then", "to", "true", "union",
    "unique", "update", "user", "using", "values", "view", "when", "where", "with",
}
ON_DELETE = {"keep": "RESTRICT", "delete": "CASCADE", "clear": "SET NULL"}
PERSONAL_KINDS = ("name", "phone", "email", "address", "birth", "id_number", "account", "health", "photo", "location", "other")


# ----------------------------------------------------------------- io helpers
def _utf8_stdio() -> None:
    """always print UTF-8 -- a Windows pipe defaults to cp1252/cp949 and Korean turns into ????"""
    for st in (sys.stdout, sys.stderr):
        try:
            st.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass


def _write(path, text: str) -> None:
    """UTF-8 with \\n line ends on every OS (Path.write_text(newline=) needs Python 3.10)"""
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)


def load(path) -> dict:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def save(model: dict, path) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(model, f, ensure_ascii=False, indent=2)
        f.write("\n")


def tables(model: dict) -> dict:
    return {t["name"]: t for t in model.get("tables", [])}


def key_columns(t: dict) -> list[dict]:
    """primary key: the id column, or ref columns marked key (1:1 child / join table)"""
    ids = [c for c in t["columns"] if c.get("type") == "id"]
    return ids or [c for c in t["columns"] if c.get("type") == "ref" and c.get("key")]


def id_column(model: dict, table: str) -> str | None:
    t = tables(model).get(table)
    if not t:
        return None
    k = key_columns(t)
    return k[0]["name"] if len(k) == 1 else None


def ordered(model: dict) -> list[dict]:
    """tables so that every referenced table comes first"""
    tb = tables(model)
    done, out = set(), []

    def visit(n, stack=()):
        if n in done or n not in tb:
            return
        if n in stack:
            return  # cycle -- check() reports it
        for c in tb[n]["columns"]:
            if c.get("type") == "ref" and c.get("ref") != n:
                visit(c.get("ref"), stack + (n,))
        done.add(n)
        out.append(tb[n])

    for t in model.get("tables", []):
        visit(t["name"])
    return out


def all_columns(t: dict) -> list[dict]:
    cols = list(t["columns"])
    if t.get("timestamps", True):
        names = {c["name"] for c in cols}
        for n, lab in (("created_at", "만든 시각"), ("updated_at", "고친 시각")):
            if n not in names:
                cols.append({"name": n, "label": lab, "type": "datetime", "required": True, "auto": True})
    return cols


# ----------------------------------------------------------------- check
def check(model: dict) -> tuple[list[str], list[str]]:
    """returns (errors, warnings) -- messages are for beginners: what is wrong + how to fix"""
    err, warn = [], []
    if model.get("format", FORMAT) != FORMAT:
        warn.append(f"format 이 '{model.get('format')}' 입니다 -- 이 도구는 '{FORMAT}' 기준으로 검사합니다")
    if model.get("db", "sqlite") not in ("sqlite", "postgresql"):
        err.append("db 는 sqlite 또는 postgresql (Supabase 는 postgresql) 이어야 합니다")
    tb = model.get("tables") or []
    if not tb:
        err.append("tables 가 비어 있습니다 -- 저장할 대상(예: 회원·주문)을 하나 이상 적어 주세요")
    seen = set()
    names = {t.get("name") for t in tb}
    for t in tb:
        tn = t.get("name", "")
        where = f"테이블 {tn or '(이름 없음)'}"
        if not NAME.match(tn or ""):
            err.append(f"{where}: 이름은 영문 소문자·숫자·밑줄(_)만, 글자로 시작 (예: order_item)")
        elif tn in RESERVED:
            err.append(f"{where}: '{tn}' 은 SQL 예약어라 쓸 수 없습니다 -- 예: {tn}s, {tn}_info")
        if tn in seen:
            err.append(f"{where}: 같은 이름의 테이블이 둘 있습니다")
        seen.add(tn)
        if not t.get("label"):
            warn.append(f"{where}: 한글 이름표(label)가 없습니다 -- 정의서·ERD 에 보일 이름을 적어 주세요")
        cols = t.get("columns") or []
        if not cols:
            err.append(f"{where}: 칸(columns)이 없습니다")
            continue
        cn_seen = set()
        ids = [c for c in cols if c.get("type") == "id"]
        keyrefs = [c for c in cols if c.get("type") == "ref" and c.get("key")]
        if len(ids) > 1:
            err.append(f"{where}: id 칸은 하나만 둘 수 있습니다")
        if not ids and not keyrefs:
            err.append(f"{where}: 행을 구별할 칸이 없습니다 -- 첫 칸에 {{\"name\": \"{tn}_id\", \"type\": \"id\"}} 를 넣으세요")
        if ids and keyrefs:
            err.append(f"{where}: id 칸이 있으면 key 표시한 ref 칸을 함께 쓸 수 없습니다")
        if ids and ids[0]["name"] != f"{tn}_id":
            warn.append(f"{where}: id 칸 이름은 '{tn}_id' 로 하면 다른 테이블에서 찾기 쉽습니다 (지금 '{ids[0]['name']}')")
        for c in cols:
            cn = c.get("name", "")
            cw = f"{tn}.{cn or '?'}"
            if not NAME.match(cn or ""):
                err.append(f"{cw}: 칸 이름은 영문 소문자·숫자·밑줄(_)만, 글자로 시작")
            elif cn in RESERVED:
                err.append(f"{cw}: '{cn}' 은 SQL 예약어 -- 예: {cn}_value, {tn}_{cn}")
            if cn in ("created_at", "updated_at") and t.get("timestamps", True):
                warn.append(f"{cw}: created_at/updated_at 은 자동으로 붙습니다 -- 지워도 됩니다")
            if cn in cn_seen:
                err.append(f"{cw}: 같은 이름의 칸이 둘 있습니다")
            cn_seen.add(cn)
            ty = c.get("type")
            if ty not in TYPES:
                err.append(f"{cw}: type '{ty}' 은 쓸 수 없습니다 -- {', '.join(TYPES)} 중 하나")
                continue
            if not c.get("label") and ty != "id":
                warn.append(f"{cw}: 한글 이름표(label)가 없습니다")
            if ty == "ref":
                r = c.get("ref")
                if r not in names:
                    err.append(f"{cw}: ref '{r}' 테이블이 없습니다")
                elif r and id_column(model, r) is None:
                    err.append(f"{cw}: '{r}' 은 칸 하나로 된 키가 없어 가리킬 수 없습니다")
                if not cn.endswith("_id"):
                    warn.append(f"{cw}: 다른 테이블을 가리키는 칸은 이름 끝을 _id 로 (예: {r}_id)")
                od = c.get("on_delete", "keep")
                if od not in ON_DELETE:
                    err.append(f"{cw}: on_delete 는 keep(지우지 못하게)·delete(같이 지움)·clear(비움) 중 하나")
                if od == "clear" and (c.get("required") or c.get("key")):
                    err.append(f"{cw}: 꼭 있어야 하는 칸(required/key)은 on_delete 'clear' 를 쓸 수 없습니다")
            elif c.get("ref"):
                err.append(f"{cw}: ref 는 type 이 'ref' 인 칸에만 씁니다")
            if ty == "owner":
                if tn.startswith("pi_"):
                    err.append(f"{cw}: 개인정보 표에는 owner 칸을 두지 않습니다 -- 짝 테이블 {tn[3:]} 에 둡니다")
                if c.get("unique") or c.get("choices") or "default" in c:
                    err.append(f"{cw}: owner 칸은 로그인한 사람이 자동으로 들어갑니다 -- unique·choices·default 를 쓰지 않습니다")
            if c.get("key") and ty != "ref":
                err.append(f"{cw}: key 표시는 ref 칸에만 (1:1 짝 테이블·연결 테이블)")
            ch = c.get("choices")
            if ch is not None:
                if not isinstance(ch, list) or not ch or not all(isinstance(x, (str, int)) for x in ch):
                    err.append(f"{cw}: choices 는 값 목록이어야 합니다 (예: [\"일반\", \"VIP\"])")
                elif ty not in ("text", "int"):
                    err.append(f"{cw}: choices 는 text·int 칸에만")
            p = c.get("personal")
            if p:
                if p is not True and p not in PERSONAL_KINDS:
                    err.append(f"{cw}: personal 은 true 또는 {', '.join(PERSONAL_KINDS)} 중 하나")
                if not tn.startswith("pi_"):
                    err.append(f"{cw}: 개인정보 칸은 따로 둡니다 -- `datalight.py split <model> --write` 로 pi_{tn} 에 옮기세요")
        owners = [c for c in cols if c.get("type") == "owner"]
        if len(owners) > 1:
            err.append(f"{where}: owner 칸은 표마다 하나만")
        acc = t.get("access")
        if acc is not None:
            if acc not in ACCESS:
                err.append(f"{where}: access 는 {', '.join(ACCESS)} 중 하나 (members=정해진 사람끼리, own=자기 행만, read=누구나 읽기만)")
            elif model.get("db", "sqlite") != "postgresql":
                warn.append(f"{where}: access 는 PostgreSQL(Supabase) 에서만 쓰입니다 -- 지금 db 는 {model.get('db', 'sqlite')}")
            elif acc == "members" and model.get("signup") != "closed":
                err.append(f"{where}: access members(로그인한 사람 모두)는 새 가입을 끈 사이트에서만 안전합니다 -- "
                           "Supabase 에서 새 가입을 끄고 정본에 \"signup\": \"closed\" 를 적거나, own 을 쓰세요")
            elif acc == "read" and tn.startswith("pi_"):
                err.append(f"{where}: 개인정보 표는 누구나 읽기(read)로 열 수 없습니다")
            elif acc == "own":
                base_t = next((x for x in tb if x.get("name") == tn[3:]), None) if tn.startswith("pi_") else t
                if not base_t or not any(c.get("type") == "owner" for c in base_t.get("columns", [])):
                    err.append(f"{where}: access own 은 " + ("짝 테이블에" if tn.startswith("pi_") else "이 표에")
                               + " owner 칸이 있어야 합니다 (예: {\"name\": \"owner_id\", \"label\": \"만든 사람\", \"type\": \"owner\"})")
        elif owners and model.get("db", "sqlite") == "postgresql":
            warn.append(f"{where}: owner 칸이 있는데 access 가 없습니다 -- 자기 행만 보게 하려면 \"access\": \"own\"")
        if tn.startswith("pi_"):
            base = tn[3:]
            if base not in names:
                warn.append(f"{where}: pi_ 테이블은 짝 테이블 '{base}' 와 1:1 로 둡니다")
            if not any(c.get("type") == "ref" and c.get("key") and c.get("ref") == base for c in cols):
                err.append(f"{where}: 짝 테이블 '{base}' 를 가리키는 key ref 칸이 필요합니다 (split 이 만들어 줍니다)")
        combos = t.get("unique") or []
        if not isinstance(combos, list) or not all(isinstance(u, list) for u in combos):
            err.append(f"{where}: 표의 unique 는 칸 묶음 목록입니다 (예: [[\"member_id\", \"dues_month\"]])")
        else:
            for u in combos:
                if len(u) < 2:
                    err.append(f"{where}: unique 묶음 {u} -- 칸 하나면 그 칸에 \"unique\": true 를 씁니다")
                miss = [x for x in u if x not in cn_seen]
                if miss:
                    err.append(f"{where}: unique 묶음에 없는 칸 {miss}")
                if len(set(u)) != len(u):
                    err.append(f"{where}: unique 묶음 {u} 에 같은 칸이 두 번")
        for i, row in enumerate(t.get("sample") or [], 1):
            bad = [k for k in row if k not in cn_seen]
            if bad:
                err.append(f"{where}: 샘플 {i}번째 행에 없는 칸 {bad}")
    # success criteria (harness layer 1): goals = questions the data must answer, not_doing = scope fence
    goals = model.get("goals")
    if goals is None:
        warn.append("성공 기준(goals)이 없습니다 -- '이 데이터로 무엇을 알 수 있어야 하나요?' 를 1~3개, 확인 SQL 과 함께 적어 두세요")
    elif not isinstance(goals, list):
        err.append("goals 는 목록이어야 합니다")
    else:
        for i, g in enumerate(goals, 1):
            gw = f"성공 기준 {i}"
            if not isinstance(g, dict) or not g.get("say") or not g.get("sql"):
                err.append(f"{gw}: say(한 문장)와 sql(확인 쿼리)이 필요합니다")
                continue
            q = g["sql"].strip().rstrip(";").strip()
            if ";" in q or not re.match(r"(?is)^(select|with)\b", q) or \
                    re.search(r"(?i)\b(insert|update|delete|drop|alter|create|attach|pragma)\b", q):
                err.append(f"{gw}: sql 은 조회(SELECT) 한 문장만 -- 데이터를 바꾸는 문장은 쓸 수 없습니다")
            ex = g.get("expect", "rows")
            if not (ex == "rows" or (isinstance(ex, int) and not isinstance(ex, bool) and ex >= 0) or isinstance(ex, list)):
                err.append(f"{gw}: expect 는 생략(1행 이상) · 행 수(숫자) · 기대 결과([[...]]) 중 하나")
    if goals and not any(t.get("sample") for t in tb):
        warn.append("성공 기준을 확인하려면 샘플 행(sample)이 필요합니다 -- 확인 쿼리가 볼 데이터를 몇 행 적어 두세요")
    nd = model.get("not_doing") or []
    if not isinstance(nd, list):
        err.append("not_doing 은 목록이어야 합니다 (예: [\"결제\", \"배송 추적\"])")
    else:
        for t in tb:
            for word in nd:
                w = str(word).strip()
                if w and (w == t.get("label") or w.lower() == str(t.get("name", "")).lower()):
                    err.append(f"테이블 {t.get('name')}: '{w}' 는 이번에 만들지 않기로 한 것(not_doing)입니다 -- 범위를 바꾸려면 사용자와 먼저 정하세요")
    if model.get("signup") not in (None, "open", "closed"):
        err.append("signup 은 open(누구나 가입) 또는 closed(새 가입을 끔)")
    outside = model.get("outside") or []
    if not isinstance(outside, list):
        err.append("outside 는 표 이름 목록입니다 (예: [\"log\"])")
    else:
        for o in outside:
            if o in names:
                err.append(f"outside 의 '{o}' 는 정본에도 있습니다 -- 둘 중 하나만")
    # cycles in required refs
    tbm = tables(model)

    def reach(a, b, seen=()):
        for c in tbm.get(a, {}).get("columns", []):
            if c.get("type") == "ref" and c.get("required", c.get("key")) and c.get("ref") in tbm:
                if c["ref"] == b or (c["ref"] not in seen and reach(c["ref"], b, seen + (a,))):
                    return True
        return False

    for t in tb:
        if t.get("name") in tbm and reach(t["name"], t["name"]):
            err.append(f"테이블 {t['name']}: 꼭 있어야 하는 참조가 돌고 돕니다(순환) -- 한 칸을 required 에서 빼 주세요")
    return err, warn


# ----------------------------------------------------------------- split
def split(model: dict) -> list[str]:
    """move personal columns of each table into pi_<table> (1:1, deleted together)"""
    done = []
    tb = tables(model)
    for t in list(model["tables"]):
        if t["name"].startswith("pi_"):
            continue
        pcols = [c for c in t["columns"] if c.get("personal")]
        if not pcols:
            continue
        idc = id_column(model, t["name"])
        if not idc:
            raise SystemExit(f"{t['name']}: 키 칸이 없어 개인정보를 나눌 수 없습니다 -- 먼저 check 오류를 고치세요")
        pn = "pi_" + t["name"]
        pt = tb.get(pn)
        if not pt:
            pt = {"name": pn, "label": (t.get("label") or t["name"]) + " 개인정보",
                  "about": f"{t.get('label') or t['name']}의 개인정보. 본 테이블과 1:1, 본 행을 지우면 같이 지워진다",
                  "columns": [{"name": idc, "label": next((c.get("label") for c in t["columns"] if c["name"] == idc), idc),
                               "type": "ref", "ref": t["name"], "key": True, "on_delete": "delete"}]}
            model["tables"].insert(model["tables"].index(t) + 1, pt)
            tb[pn] = pt
        for c in pcols:
            t["columns"].remove(c)
            pt["columns"].append(c)
        # samples follow their columns
        if t.get("sample"):
            moved = [c["name"] for c in pcols]
            prow = []
            for i, row in enumerate(t["sample"], 1):
                pr = {k: row.pop(k) for k in moved if k in row}
                if pr:
                    pr[idc] = row.get(idc, i)
                    prow.append(pr)
            if prow:
                pt.setdefault("sample", []).extend(prow)
        done.append(f"{t['name']} → {pn}: {', '.join(c['name'] for c in pcols)}")
    return done


# ----------------------------------------------------------------- build
def _q(v) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def sql_type(c: dict, db: str) -> str:
    ty = c["type"]
    if db == "postgresql":
        return {"id": "bigint", "ref": "bigint", "text": "text", "long_text": "text", "int": "integer",
                "number": "numeric", "money": "numeric(15,2)", "bool": "boolean", "date": "date",
                "datetime": "timestamptz", "owner": "uuid"}[ty]
    return {"id": "INTEGER", "ref": "INTEGER", "text": "TEXT", "long_text": "TEXT", "int": "INTEGER", "number": "REAL",
            "money": "NUMERIC", "bool": "INTEGER", "date": "TEXT", "datetime": "TEXT", "owner": "TEXT"}[ty]


def ddl(model: dict, db: str | None = None, only: set | None = None) -> str:
    db = db or model.get("db", "sqlite")
    pg = db == "postgresql"
    out = [f"-- {model.get('project', '')} -- {db} schema (data-light)",
           "-- 이 파일은 model.light.json 에서 만들어졌습니다. 직접 고치지 말고 정본을 고친 뒤 다시 build 하세요."]
    if not pg:
        out.append("PRAGMA foreign_keys = ON;  -- SQLite 는 이 줄이 있어야 참조 규칙이 동작합니다")
    else:
        if any(c.get("type") == "owner" for t in model.get("tables", []) for c in t["columns"]):
            out.append("-- owner 칸이 auth.uid()(로그인한 사람)를 쓰므로 Supabase 용입니다.")
        out.append("BEGIN;")
    for t in ordered(model):
        tn = t["name"]
        if only is not None and tn not in only:
            continue
        lines = []
        keys = key_columns(t)
        for c in all_columns(t):
            ty = c["type"]
            s = f"  {c['name']} {sql_type(c, db)}"
            if ty == "id":
                s += " GENERATED ALWAYS AS IDENTITY PRIMARY KEY" if pg else " PRIMARY KEY AUTOINCREMENT"
            elif ty == "owner":
                # Supabase: the signed-in user. SQLite has no sign-in, so it stays a plain optional text there
                s += " NOT NULL DEFAULT auth.uid()" if pg else ""
            else:
                if c.get("required") or c.get("key"):
                    s += " NOT NULL"
                if c.get("auto"):
                    s += " DEFAULT " + ("now()" if pg else "CURRENT_TIMESTAMP")
                elif "default" in c:
                    s += " DEFAULT " + (_q(c["default"]) if not (ty == "bool" and not pg) else str(int(bool(c["default"]))))
                if c.get("unique"):
                    s += " UNIQUE"
                if ty == "ref":
                    rid = id_column(model, c["ref"])
                    s += f" REFERENCES {c['ref']} ({rid}) ON DELETE {ON_DELETE[c.get('on_delete', 'keep')]}"
                if c.get("choices"):
                    s += f" CHECK ({c['name']} IN ({', '.join(_q(x) for x in c['choices'])}))"
                elif ty == "bool" and not pg:
                    s += f" CHECK ({c['name']} IN (0, 1))"
                elif ty == "money":
                    s += f" CHECK ({c['name']} >= 0)"
            lines.append(s)
        if keys and keys[0]["type"] == "ref":
            lines.append(f"  PRIMARY KEY ({', '.join(k['name'] for k in keys)})")
        for u in t.get("unique") or []:
            lines.append(f"  UNIQUE ({', '.join(u)})")
        head = f"-- {t.get('label', tn)}" + (f": {t['about']}" if t.get("about") else "")
        out.append(f"\n{head}\nCREATE TABLE {tn} (\n" + ",\n".join(lines) + "\n);")
        for c in t["columns"]:
            if c["type"] == "ref" and not (c.get("key") and len(keys) == 1):
                out.append(f"CREATE INDEX ix_{tn}_{c['name']} ON {tn} ({c['name']});")
    if pg:
        out.append("")
        for t in ordered(model):
            out.append(f"COMMENT ON TABLE {t['name']} IS {_q(t.get('label', t['name']) + (' -- ' + t['about'] if t.get('about') else ''))};")
            for c in all_columns(t):
                lab = (c.get("label") or c["name"]) + (" [개인정보]" if c.get("personal") else "")
                out.append(f"COMMENT ON COLUMN {t['name']}.{c['name']} IS {_q(lab)};")
        out += rls_sql(model, only)
        out.append("COMMIT;")
    return "\n".join(out) + "\n"


def _policies(model: dict, t: dict) -> list[str]:
    """real policies from the model's "access" (members / own / read)"""
    tn, acc = t["name"], t.get("access")
    if acc == "members":
        return [f"CREATE POLICY {tn}_members ON {tn} FOR ALL TO authenticated USING (true) WITH CHECK (true);"]
    if acc == "read":
        return [f"CREATE POLICY {tn}_read ON {tn} FOR SELECT TO anon, authenticated USING (true);"]
    if acc == "own":
        own = next((c["name"] for c in t["columns"] if c["type"] == "owner"), None)
        if own:
            cond = f"{own} = auth.uid()"
        else:                                   # pi_ table: owner lives on the paired table
            base = tables(model)[tn[3:]]
            bo = next(c["name"] for c in base["columns"] if c["type"] == "owner")
            k = id_column(model, base["name"])
            cond = f"EXISTS (SELECT 1 FROM {base['name']} b WHERE b.{k} = {tn}.{k} AND b.{bo} = auth.uid())"
        return [f"CREATE POLICY {tn}_own ON {tn} FOR ALL TO authenticated USING ({cond}) WITH CHECK ({cond});"]
    return []


def rls_sql(model: dict, only=None) -> list[str]:
    """Supabase: the public (anon / publishable) key in the site can reach every table, so RLS goes on for all.
    No policy = nobody can read or write through that key (safe default)."""
    out = ["", "-- ===== 행 접근 규칙 (RLS) =====",
           "-- Supabase 에서는 사이트 코드에 넣는 공개용 키(anon · publishable)로 누구나 모든 표에 접근을 시도할 수 있습니다.",
           "-- 그래서 모든 표에 RLS 를 켭니다. 규칙(policy)이 없는 표는 공개용 키로 아무것도 읽고 쓸 수 없습니다 (안전한 기본값).",
           "-- 규칙은 정본의 표에 \"access\" 를 적고 다시 build 하면 아래에 만들어집니다 (권장: members · own · read).",
           "-- 틀만 있는 표는 이 파일을 고치지 말고, SQL 편집기에 붙여 넣은 뒤 고른 틀의 주석을 풀어 실행하고,",
           "-- 실행한 규칙은 data/supabase-policies.sql 에 적어 둡니다 (이 파일을 고치면 verify 가 '정본과 다릅니다'로 멈춥니다).",
           "-- 사이트 코드에는 공개용 키만. service_role · sb_secret_ 키는 절대 넣지 않습니다. 적용 뒤 Supabase 의 Security Advisor 를 확인하세요.",
           "-- (일반 PostgreSQL 에서는 표 주인 계정에는 RLS 가 적용되지 않아 동작이 바뀌지 않습니다.)"]
    if model.get("signup") == "closed":
        out.append("-- 이 사이트는 새 가입을 끈 것으로 정했습니다(signup: closed). Supabase 의 가입 설정이 꺼져 있는지 확인하세요.")
    for t in ordered(model):
        tn = t["name"]
        if only is not None and tn not in only:
            continue
        out.append(f"ALTER TABLE {tn} ENABLE ROW LEVEL SECURITY;")
        real = _policies(model, t)
        if real:
            out += real
            continue
        members = (f"--   [정해진 사람끼리 · 새 가입을 끈 경우에만] CREATE POLICY {tn}_members ON {tn} FOR ALL TO authenticated"
                   " USING (true) WITH CHECK (true);  -- 가입이 켜져 있으면 아무나 가입해 읽고 씁니다")
        own = (f"--   [가입형·자기 행] CREATE POLICY {tn}_own ON {tn} FOR ALL TO authenticated"
               " USING (owner_id = auth.uid()) WITH CHECK (owner_id = auth.uid());  -- 정본에 owner 칸 + \"access\": \"own\" 이 더 쉽습니다")
        if tn.startswith("pi_"):
            out += [f"-- {tn} 는 개인정보: 누구나 읽기 규칙은 쓰지 않습니다. 본인 · 관리자 · (가입을 끈 경우) 정해진 사람만.",
                    own, members,
                    "--   [관리자만]     관리자를 구별하는 방법(관리자 표 또는 JWT claim)을 정한 뒤 USING 에 그 조건을 씁니다"]
        else:
            out += [members, own,
                    f"--   [누구나 읽기만]   CREATE POLICY {tn}_read ON {tn} FOR SELECT TO anon, authenticated USING (true);"]
    return out


def erd(model: dict) -> str:
    out = [f"# {model.get('project', '')} -- ERD", "", "```mermaid", "erDiagram"]
    for t in ordered(model):
        keys = {k["name"] for k in key_columns(t)}
        out.append(f"    {t['name']} {{")
        for c in all_columns(t):
            if c.get("auto"):
                continue
            mark = " PK" if c["name"] in keys else ""
            mark += (" FK" if mark == "" else ",FK") if c["type"] == "ref" else ""
            label = (c.get("label") or "").replace('"', "'") + (" (개인정보)" if c.get("personal") else "")
            out.append(f"        {c['type']} {c['name']}{mark} \"{label}\"")
        out.append("    }")
    for t in ordered(model):
        keys = key_columns(t)
        for c in t["columns"]:
            if c["type"] != "ref":
                continue
            one_to_one = c.get("key") and len(keys) == 1
            left = "||" if c.get("required") or c.get("key") else "|o"
            right = "||" if one_to_one else "o{"
            out.append(f"    {c['ref']} {left}--{right} {t['name']} : \"{(c.get('label') or c['name']).replace(chr(34), '')}\"")
    out += ["```", ""]
    return "\n".join(out)


TYPE_KO = {"id": "번호(자동)", "ref": "다른 표 가리킴", "owner": "만든 사람(로그인)", "text": "글자", "long_text": "긴 글", "int": "정수", "number": "숫자",
           "money": "금액", "bool": "예/아니오", "date": "날짜", "datetime": "날짜+시각"}


def spec_rows(model: dict):
    for t in ordered(model):
        keys = {k["name"] for k in key_columns(t)}
        for i, c in enumerate(all_columns(t), 1):
            rule = []
            if c.get("required") or c["name"] in keys:
                rule.append("필수")
            if c.get("unique"):
                rule.append("중복 금지")
            if c.get("choices"):
                rule.append("선택: " + "/".join(map(str, c["choices"])))
            if c["type"] == "ref":
                rule.append(f"→ {c['ref']} ({ {'keep': '지우지 못하게', 'delete': '같이 지움', 'clear': '비움'}[c.get('on_delete', 'keep')] })")
            yield (t["name"], t.get("label", ""), i, c["name"], c.get("label", ""), TYPE_KO[c["type"]],
                   "키" if c["name"] in keys else "", ", ".join(rule), str(c.get("personal") or "") if c.get("personal") else "",
                   c.get("about", "") or ("자동" if c.get("auto") else ""))


SPEC_HEAD = ("테이블", "테이블 이름표", "순번", "칸", "칸 이름표", "자료형", "키", "규칙", "개인정보", "설명")


def spec_md(model: dict) -> str:
    out = [f"# {model.get('project', '')} -- 테이블 정의서", "", f"DB: {model.get('db', 'sqlite')}", ""]
    if model.get("goals") or model.get("not_doing"):
        out += ["## 성공 기준 -- 이 데이터로 할 수 있어야 하는 것", "",
                "`datalight.py verify` 가 아래 확인 쿼리를 샘플 데이터로 실제로 돌려 봅니다.", ""]
        out += [f"{i}. {g.get('say')}  \n   `{g.get('sql')}`" for i, g in enumerate(model.get("goals") or [], 1)]
        if model.get("not_doing"):
            out += ["", "**이번에 만들지 않을 것:** " + ", ".join(map(str, model["not_doing"]))]
        out.append("")
    cur = None
    for r in spec_rows(model):
        if r[0] != cur:
            cur = r[0]
            t = tables(model)[cur]
            combo = "".join(f"\n\n함께 중복 금지: {' + '.join(u)}" for u in t.get("unique") or [])
            out += ["", f"## {r[1] or cur} (`{cur}`)", "", t.get("about", "") + combo, "",
                    "| " + " | ".join(SPEC_HEAD[2:]) + " |", "|" + "---|" * (len(SPEC_HEAD) - 2)]
        out.append("| " + " | ".join(str(x).replace("|", "/") for x in r[2:]) + " |")
    return "\n".join(out) + "\n"


def spec_xlsx(model: dict, path) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    ws = wb.active
    ws.title = "정의서"
    ws.append(SPEC_HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E4E7F7")
    pink = PatternFill("solid", fgColor="FBE0E5")
    for r in spec_rows(model):
        ws.append(r)
        if r[8]:
            for c in ws[ws.max_row]:
                c.fill = pink
    ws.freeze_panes = "A2"
    for col, wdt in zip("ABCDEFGHIJ", (18, 16, 6, 18, 16, 12, 6, 36, 10, 40)):
        ws.column_dimensions[col].width = wdt
    wb.save(path)
    return True


def build(model: dict, out_dir, db: str | None = None) -> list[str]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    files = []
    for name, text in (("schema.sql", ddl(model, db)), ("erd.md", erd(model)), ("spec.md", spec_md(model))):
        _write(out / name, text)
        files.append(name)
    if spec_xlsx(model, out / "spec.xlsx"):
        files.append("spec.xlsx")
    return files


# ----------------------------------------------------------------- try (SQLite probe)
def _fill(c: dict, n: int):
    ty = c["type"]
    if c.get("choices"):
        return c["choices"][0]
    return {"text": f"샘플{n}", "long_text": f"샘플 글 {n}", "int": n, "number": n + 0.5, "money": 1000 * n,
            "bool": 1, "date": "2026-01-0%d" % min(n, 9), "datetime": "2026-01-01 09:00:00"}.get(ty, f"v{n}")


def _norm(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return float(v)
    return v


def run_goals(con, model: dict) -> tuple[bool, list[str]]:
    ok, lines = True, []
    for i, g in enumerate(model.get("goals") or [], 1):
        label = f"성공 기준 {i} '{g.get('say')}'"
        try:
            rows = [list(r) for r in con.execute(g["sql"].strip().rstrip(";")).fetchall()]
        except sqlite3.Error as e:
            ok = False
            lines.append(f"FAIL: {label} -- 쿼리 실패: {e}")
            continue
        ex = g.get("expect", "rows")
        if ex == "rows":
            good, why = bool(rows), "결과가 없습니다"
        elif isinstance(ex, int):
            good, why = len(rows) == ex, f"{ex}행이어야 하는데 {len(rows)}행"
        else:
            key = lambda r: json.dumps([_norm(x) for x in r], ensure_ascii=False)  # noqa: E731
            good = sorted(map(key, rows)) == sorted(map(key, ex))
            why = f"기대 {ex} / 실제 {rows}"
        lines.append(f"OK: {label}" if good else f"FAIL: {label} -- {why}")
        ok &= good
    return ok, lines


def try_sqlite(model: dict, ddl_text: str | None = None) -> tuple[bool, list[str]]:
    """create the SQLite schema in memory, insert samples, run the goals, then check each rule really blocks bad rows"""
    con = sqlite3.connect(":memory:")
    con.execute("PRAGMA foreign_keys = ON")
    lines, ok = [], True
    try:
        con.executescript(ddl_text if ddl_text is not None else ddl(model, "sqlite"))
        lines.append(f"OK: 테이블 {len(model['tables'])}개 만들기")
    except sqlite3.Error as e:
        con.close()
        return False, [f"FAIL: 스키마 만들기 실패 -- {e}"]
    first_ids: dict[str, int] = {}

    def insert(t, row):
        cols = list(row)
        cur = con.execute(f"INSERT INTO {t['name']} ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                          [int(v) if isinstance(v, bool) else v for v in row.values()])
        return cur.lastrowid

    def minimal_row(t, n=1):
        row = {}
        for c in t["columns"]:
            if c["type"] == "id":
                continue
            if c["type"] == "ref":
                if c.get("required") or c.get("key"):
                    row[c["name"]] = first_ids.get(c["ref"])
                continue
            if c.get("required") or c.get("unique"):
                row[c["name"]] = _fill(c, n)
        return row

    for t in ordered(model):
        rows = t.get("sample") or []
        try:
            for r in rows:
                rid = insert(t, r)
                first_ids.setdefault(t["name"], r.get(id_column(model, t["name"]) or "", rid) or rid)
            if not rows:
                rid = insert(t, minimal_row(t))
                kc = key_columns(t)
                first_ids.setdefault(t["name"], rid if kc and kc[0]["type"] == "id" else first_ids.get(kc[0]["ref"]) if kc else rid)
            lines.append(f"OK: {t['name']} 샘플 {len(rows) or 1}행 넣기")
        except sqlite3.Error as e:
            ok = False
            lines.append(f"FAIL: {t['name']} 샘플 넣기 실패 -- {e}")
    con.commit()
    gok, glines = run_goals(con, model)
    ok &= gok
    lines += glines

    def expect_block(label, fn):
        nonlocal ok
        con.execute("SAVEPOINT p")
        try:
            fn()
            ok = False
            lines.append(f"FAIL: {label} -- 막혀야 하는데 들어갔습니다")
        except sqlite3.Error:
            lines.append(f"OK: {label} 막힘")
        con.execute("ROLLBACK TO p")
        con.execute("RELEASE p")

    for t in ordered(model):
        base = minimal_row(t, 7)
        for c in t["columns"]:
            n = f"{t['name']}.{c['name']}"
            if c["type"] in ("id",) or (c["type"] == "ref" and c.get("key") and len(key_columns(t)) == 1):
                continue
            if c.get("required") and c["type"] != "ref":
                expect_block(f"{n} 비우기", lambda c=c: insert(t, {**base, c["name"]: None}))
            if c.get("choices"):
                expect_block(f"{n} 목록 밖 값", lambda c=c: insert(t, {**base, c["name"]: "__없는값__"}))
            if c["type"] == "money":
                expect_block(f"{n} 음수 금액", lambda c=c: insert(t, {**base, c["name"]: -1}))
            if c["type"] == "ref":
                expect_block(f"{n} 없는 대상 가리키기", lambda c=c: insert(t, {**base, c["name"]: 987654321}))
                if c.get("required"):
                    expect_block(f"{n} 비우기", lambda c=c: insert(t, {**base, c["name"]: None}))
        uq = [c for c in t["columns"] if c.get("unique")]
        for c in uq:
            def dup(c=c):
                r1 = {**base, c["name"]: _fill(c, 99)}
                for o in uq:
                    if o is not c:
                        r1[o["name"]] = _fill(o, 98)
                insert(t, r1)
                r2 = dict(r1)
                for o in uq:
                    if o is not c:
                        r2[o["name"]] = _fill(o, 97)
                insert(t, r2)
            expect_block(f"{t['name']}.{c['name']} 중복", dup)
        for u in t.get("unique") or []:
            cols_by = {c["name"]: c for c in t["columns"]}

            def dup_combo(u=u, cols_by=cols_by, uq=uq, base=base, t=t):
                r1 = dict(base)
                for x in u:
                    c = cols_by[x]
                    r1[x] = first_ids.get(c["ref"]) if c["type"] == "ref" else _fill(c, 55)
                for o in uq:
                    if o["name"] not in u:
                        r1[o["name"]] = _fill(o, 61)
                insert(t, r1)
                r2 = dict(r1)
                for o in uq:
                    if o["name"] not in u:
                        r2[o["name"]] = _fill(o, 62)
                insert(t, r2)
            expect_block(f"{t['name']} ({' + '.join(u)}) 함께 중복", dup_combo)
    # deleting a parent that children point to with on_delete keep
    for t in ordered(model):
        for c in t["columns"]:
            if c["type"] == "ref" and c.get("on_delete", "keep") == "keep" and c.get("ref") in first_ids:
                if con.execute(f"SELECT count(*) FROM {t['name']} WHERE {c['name']} = ?", (first_ids[c['ref']],)).fetchone()[0]:
                    rid = id_column(model, c["ref"])
                    expect_block(f"{c['ref']} 삭제(→ {t['name']} 가 가리키는 중)",
                                 lambda c=c, rid=rid: con.execute(f"DELETE FROM {c['ref']} WHERE {rid} = ?", (first_ids[c['ref']],)))
    con.close()
    return ok, lines


# ----------------------------------------------------------------- verify (harness gate)
def _mutants(model: dict, base: str):
    """one planted bug per rule kind, aimed at a column the model says is protected"""
    def first(pred):
        for t in ordered(model):
            for c in t["columns"]:
                if pred(t, c):
                    return t["name"], c["name"]
        return None

    def cut(target, rx_tail, keep_group=False):
        if not target:
            return None
        t, c = target
        start = base.index(f"CREATE TABLE {t} (")
        end = base.index("\n);", start)
        m = re.compile(rf"(\n  {c} \w+(?:\(\d+,\d+\))?[^\n]*?)" + rx_tail).search(base, start, end)
        if not m:
            return None
        return base[:m.start()] + m.group(1) + base[m.end():]

    single_key = lambda t, c: c.get("key") and len(key_columns(t)) == 1  # noqa: E731
    yield "필수값(NOT NULL) 하나 빼기", cut(first(lambda t, c: c.get("required") and c["type"] not in ("id", "ref")), " NOT NULL")
    yield "선택값 규칙(CHECK IN) 하나 빼기", cut(first(lambda t, c: c.get("choices")), r" CHECK \(\w+ IN \([^)]*\)\)")
    yield "금액 음수 금지 하나 빼기", cut(first(lambda t, c: c["type"] == "money"), r" CHECK \(\w+ >= 0\)")
    yield "중복 금지(UNIQUE) 하나 빼기", cut(first(lambda t, c: c.get("unique") and c["type"] != "id"), " UNIQUE")
    combo = next(((t["name"], u) for t in ordered(model) for u in (t.get("unique") or [])), None)
    if combo:
        line = f",\n  UNIQUE ({', '.join(combo[1])})"
        yield "함께 중복 금지(UNIQUE 묶음) 하나 빼기", base.replace(line, "", 1) if line in base else None
    yield "참조(REFERENCES) 하나 빼기", cut(first(lambda t, c: c["type"] == "ref" and not single_key(t, c)),
                                       r" REFERENCES \w+ \(\w+\) ON DELETE (?:RESTRICT|CASCADE|SET NULL)")


def mutation_check(model: dict) -> list[tuple[str, bool]]:
    """plant one bug at a time in the SQL and confirm `try` notices (does the safety device really catch?)"""
    base = ddl(model, "sqlite")
    out = []
    for label, broken in _mutants(model, base):
        if broken is None:
            continue
        ok, _ = try_sqlite(model, broken)
        out.append((label, not ok))
    return out


def stale_outputs(model: dict, out_dir) -> list[str]:
    """files in data/out that no longer match the model (someone edited the model but did not rebuild)"""
    out = Path(out_dir)
    if not out.is_dir():
        return []
    want = {"schema.sql": None, "erd.md": erd(model), "spec.md": spec_md(model)}
    stale = []
    for name, text in want.items():
        f = out / name
        if not f.exists():
            continue
        if name == "schema.sql":
            first = f.read_text(encoding="utf-8").split("\n", 1)[0]
            db = "postgresql" if "postgresql" in first else "sqlite"
            text = ddl(model, db)
        if f.read_text(encoding="utf-8") != text:
            stale.append(name)
    return stale


def db_drift(model: dict, db_path) -> tuple[list[str], list[str]]:
    """compare a real SQLite DB with the model. (irreversible, notes)
    irreversible = tables/columns that hold data in the DB but are gone from the model"""
    con = sqlite3.connect(f"file:{Path(db_path).as_posix()}?mode=ro", uri=True)
    bad, notes = [], []
    try:
        have = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}
        tb = tables(model)
        outside = set(model.get("outside") or [])
        for t in sorted(have):
            if t in outside:
                notes.append(f"{t} 는 정본 밖(outside)으로 둔 표 -- 검사하지 않습니다")
                continue
            n = con.execute(f'SELECT count(*) FROM "{t}"').fetchone()[0]
            if t not in tb:
                (bad if n else notes).append(f"DB 의 {t} 가 정본에서 빠졌습니다 ({n}행)" + (" -- 지우면 되돌릴 수 없습니다" if n else ""))
                continue
            cols = [r[1] for r in con.execute(f'PRAGMA table_info("{t}")')]
            want = {c["name"] for c in all_columns(tb[t])}
            for c in cols:
                if c not in want:
                    filled = con.execute(f'SELECT count(*) FROM "{t}" WHERE "{c}" IS NOT NULL').fetchone()[0]
                    (bad if filled else notes).append(f"DB 의 {t}.{c} 가 정본에서 빠졌습니다 (값 {filled}개)"
                                                      + (" -- 지우면 되돌릴 수 없습니다" if filled else ""))
            for c in sorted(want - set(cols)):
                notes.append(f"{t}.{c} 는 정본에만 있습니다 (DB 에 아직 반영 전)")
        for t in sorted(set(tb) - have):
            notes.append(f"{t} 는 정본에만 있습니다 (DB 에 아직 반영 전)")
    finally:
        con.close()
    return bad, notes


def verify(model_path, db_path=None, use_db=True, mutate=False) -> tuple[bool, list[str]]:
    """the single pass/fail gate. Used by people, by Claude before saying 'done', by the commit hook and by CI"""
    m = load(model_path)
    lines, ok = [], True
    errs, warns = check(m)
    lines += [f"FAIL: {e}" for e in errs] + [f"참고: {w}" for w in warns]
    if errs:
        return False, lines + ["결과: 설계 검사(check) 실패 -- 나머지 검사는 건너뜁니다"]
    lines.append("OK: 설계 검사")
    tok, tlines = try_sqlite(m)
    ok &= tok
    lines += [l for l in tlines if l.startswith("FAIL") or "성공 기준" in l]
    if tok:
        lines.append("OK: 규칙 확인 (나쁜 값이 모두 막힘)")
    out = Path(model_path).parent / "out"
    stale = stale_outputs(m, out)
    if stale:
        ok = False
        lines.append(f"FAIL: {out} 의 {', '.join(stale)} 가 정본과 다릅니다 -- `build` 로 다시 만드세요")
    if use_db:
        dbp = Path(db_path) if db_path else Path(model_path).parent / "app.db"
        if dbp.exists():
            bad, notes = db_drift(m, dbp)
            ok &= not bad
            lines += [f"FAIL: {b}" for b in bad] + [f"참고: {n}" for n in notes]
            if not bad:
                lines.append(f"OK: {dbp.name} 와 비교 -- 데이터가 사라질 변경 없음")
    if mutate:
        res = mutation_check(m)
        for label, caught in res:
            lines.append(("OK: 장치 검증 -- " if caught else "FAIL: 장치 검증 -- ") + label + (" → 잡음" if caught else " → 못 잡음 (검사에 구멍)"))
            ok &= caught
    lines.append("결과: 통과" if ok else "결과: 실패 -- FAIL 줄을 고친 뒤 다시 verify")
    return ok, lines


HOOK_MARK = "# data-light pre-commit"
HOOK_END = "# data-light end"
HOOK_VER = "# data-light-version: "
HOOK = HOOK_MARK + """ (datalight.py hook install)
""" + HOOK_VER + __version__ + """
# 정본(model.light.json)이나 그 out/ 이 바뀐 커밋은 verify 를 통과해야 한다.
# 멈추면 FAIL 줄을 고친 뒤 다시 저장(커밋)하세요. 검사를 건너뛰지 않습니다.
# (괄호 안에서만 끝내므로 이 아래에 다른 훅 내용이 있어도 그대로 돕니다)
(
  models=$(git -c core.quotepath=false diff --cached --name-only --diff-filter=ACMR | while IFS= read -r f; do
    case "$f" in .claude/skills/*|.agents/skills/*) continue ;; esac   # 설치된 스킬의 예제는 검사하지 않음
    case "$f" in
      *model.light.json) printf '%s\\n' "$f" ;;
      */out/*) d=${f%/out/*}; [ -f "$d/model.light.json" ] && printf '%s\\n' "$d/model.light.json" ;;
      out/*) [ -f model.light.json ] && printf '%s\\n' model.light.json ;;
    esac
  done | sort -u)
  [ -z "$models" ] && exit 0
  PY=""
  for c in python3 python py; do "$c" -c "import sys" >/dev/null 2>&1 && { PY=$c; break; }; done  # Windows: python3 can be a Store stub
  [ -z "$PY" ] && { echo "data-light: python 을 찾지 못해 검사를 건너뜁니다"; exit 0; }
  printf '%s\\n' "$models" | while IFS= read -r m; do
    "$PY" .data-light/datalight.py verify "$m" || { echo "data-light: $m 검사 실패 -- 커밋을 멈췄습니다"; exit 1; }
  done
) || exit 1
""" + HOOK_END + "\n"
CI = """name: data-light verify
on: [push, pull_request]
jobs:
  verify:
    runs-on: ubuntu-latest
    steps:
      - uses: actions/checkout@v5
      - uses: actions/setup-python@v6
        with: { python-version: "3.12" }
      - name: verify every model.light.json (installed skill examples excluded)
        run: |
          git -c core.quotepath=false ls-files '*model.light.json' | grep -v -E '^(\\.claude|\\.agents)/skills/' > "$RUNNER_TEMP/models.txt" || true
          while IFS= read -r m; do python .data-light/datalight.py verify "$m" --no-db || exit 1; done < "$RUNNER_TEMP/models.txt"
"""


GITIGNORE_LINES = ("data/app.db", "data/backup/", "data/*.csv", "data/*.xlsx", "data/out/spec.xlsx",
                   "data/import_plan*.json", "import_plan*.json", ".env")


def _git(repo: Path, *args) -> str:
    import subprocess
    try:
        r = subprocess.run(["git", "-C", str(repo), *args], capture_output=True, text=True, encoding="utf-8", errors="replace")
        return r.stdout.strip() if r.returncode == 0 else ""
    except OSError:
        return ""


def _hooks_dir(repo: Path) -> tuple[Path, str]:
    """where git really runs hooks: core.hooksPath (husky, lefthook, shared folder) or the (worktree's common) .git/hooks"""
    gp = _git(repo, "rev-parse", "--git-path", "hooks")
    d = Path(gp) if gp else repo / ".git" / "hooks"
    if not d.is_absolute():
        d = repo / d
    try:
        shown = Path(os.path.relpath(d, repo)).as_posix()
    except ValueError:
        shown = d.as_posix()
    return d, shown


def _hook_label(repo: Path, shown: str) -> str:
    return f"{shown}/pre-commit" + (" (git 설정 core.hooksPath)" if _git(repo, "config", "--get", "core.hooksPath") else "")


def _settings(repo: Path) -> dict:
    f = repo / ".data-light" / "settings.json"
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def hook_status(repo=".") -> tuple[bool, list[str]]:
    repo = Path(repo).resolve()
    hd, shown = _hooks_dir(repo)
    where = _hook_label(repo, shown)
    hook = hd / "pre-commit"
    tool = repo / ".data-light" / "datalight.py"
    lines, same = [f"이 도구: data-light {__version__}"], True
    if not hook.exists() or HOOK_MARK not in hook.read_text(encoding="utf-8", errors="replace"):
        return False, lines + [f"훅: 없음 ({where}) -- `hook install` 필요"]
    m = re.search(re.escape(HOOK_VER) + r"(\S+)", hook.read_text(encoding="utf-8", errors="replace"))
    hv = m.group(1) if m else "(0.2 이하)"
    lines.append(f"훅: {where} · {hv}")
    same &= hv == __version__
    if tool.exists():
        tm = re.search(r'__version__ = "([^"]+)"', tool.read_text(encoding="utf-8", errors="replace"))
        tv = tm.group(1) if tm else "(0.2 이하)"
        lines.append(f"도구 사본: .data-light/datalight.py · {tv}")
        same &= tv == __version__
    else:
        same = False
        lines.append("도구 사본: 없음")
    lines.append("결과: 최신" if same else "결과: 버전이 다릅니다 -- `hook install` 로 다시 설치")
    return same, lines


def hook_install(repo=".", ci=None, append=False) -> list[str]:
    """ci: True/False/None(=remembered choice, default on). append: add our block to someone else's pre-commit"""
    repo = Path(repo).resolve()
    if _git(repo, "rev-parse", "--is-inside-work-tree") != "true":
        raise SystemExit(f"{repo} 는 git 저장소가 아닙니다 (git init 먼저)")
    done = []
    tool = repo / ".data-light" / "datalight.py"
    tool.parent.mkdir(exist_ok=True)
    tool.write_bytes(Path(__file__).read_bytes())
    done.append(f".data-light/datalight.py {__version__} (도구 사본 -- 커밋해 두면 CI 와 다른 사람도 같은 검사)")
    st = _settings(repo)
    if ci is None:
        ci = st.get("ci", True)
    if st.get("ci", True) != ci:
        st["ci"] = ci
        _write(repo / ".data-light" / "settings.json", json.dumps(st, ensure_ascii=False, indent=2) + "\n")
    hd, shown = _hooks_dir(repo)
    where = _hook_label(repo, shown)
    hook = hd / "pre-commit"
    hd.mkdir(parents=True, exist_ok=True)
    old = hook.read_text(encoding="utf-8", errors="replace") if hook.exists() else ""
    if HOOK_MARK in old:                       # ours (alone or appended) -> replace only our block, keep before and after
        i = old.index(HOOK_MARK)
        j = old.find(HOOK_END, i)
        rest = old[j + len(HOOK_END):].lstrip("\n") if j >= 0 else ""      # 0.3 hooks had no end mark: ours ran to the end
        _write(hook, old[:i] + HOOK + (("\n" + rest) if rest else ""))
        done.append(f"{where} (다시 설치 · {__version__})")
    elif old.strip() and not append:
        done.append(f"{where} 이 이미 있어 건드리지 않았습니다 -- 끝에 덧붙이려면 `hook install --append`")
    elif old.strip():
        _write(hook, old.rstrip("\n") + "\n\n" + HOOK)
        done.append(f"{where} 끝에 덧붙임 -- 그 위에 `exit 0` 같은 줄이 있으면 여기까지 오지 않으니 확인하세요")
    else:
        _write(hook, "#!/bin/sh\n" + HOOK)
        done.append(f"{where} (커밋할 때 자동 검사)")
    try:
        hook.chmod(0o755)
    except OSError:
        pass
    wf = repo / ".github" / "workflows" / "data-light.yml"
    if not ci:
        done.append("CI 파일은 만들지 않음 (--no-ci 를 기억해 두었습니다. 다시 켜려면 `hook install --ci`)")
    elif not wf.exists():
        wf.parent.mkdir(parents=True, exist_ok=True)
        _write(wf, CI)
        done.append(".github/workflows/data-light.yml (GitHub 에 올릴 때 자동 검사)")
    gi = repo / ".gitignore"
    have = gi.read_text(encoding="utf-8").splitlines() if gi.exists() else []
    add = [x for x in GITIGNORE_LINES if x not in have]
    if add:
        with open(gi, "a", encoding="utf-8") as f:
            f.write(("\n" if have and have[-1] else "") + "# data-light: 실제 데이터·가져오기 계획·비밀 값은 올리지 않는다\n"
                    + "\n".join(add) + "\n")
        done.append(".gitignore 에 " + ", ".join(add))
    return done


# ----------------------------------------------------------------- import
PI_HINT = [("name", r"이름|성명|성함|^name$|full_?name|(회원|고객|학생|사원|직원|담당자|신청자|보호자|환자|학부모|손님)명|대표자"), ("phone", r"전화|휴대|핸드폰|연락처|phone|mobile|tel"),
           ("email", r"이메일|메일|e-?mail"), ("address", r"주소|address|addr"), ("birth", r"생년|생일|birth"),
           ("id_number", r"주민|여권|면허번호|passport"), ("account", r"계좌|카드번호|account_no|card")]
NOT_PI = re.compile(r"상호|가게|업체|회사|기관|시설|상품|제품|도로명|소재지|대표전화|팩스")
# English column names in an existing DB (import schema) -- person words only, so dish_name / hotel are not flagged
SCHEMA_PI = [("name", r"(customer|member|user|student|employee|staff|patient|parent|guest|contact|full|person|real)_?name|^name$"),
             ("phone", r"phone|mobile|(^|_)tel($|_)"), ("email", r"e_?mail"), ("address", r"addr"),
             ("birth", r"birth|(^|_)dob($|_)"), ("id_number", r"ssn|passport|resident_?(no|num)"),
             ("account", r"account_?no|card_?no|bank_?account")]


def _read_table(path: str, sheet=None) -> tuple[list[str], list[list]]:
    p = path.lower()
    if p.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("엑셀 파일은 openpyxl 이 필요합니다: pip install openpyxl (또는 csv 로 저장해서 넣으세요)")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows = [[("" if v is None else v) for v in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        rows = [r for r in rows if any(str(v).strip() for v in r)]
        return [str(h).strip() for h in rows[0]], rows[1:]
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:20000], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(x.strip() for x in r)]
    return [h.strip() for h in rows[0]], rows[1:]


def _infer(values: list) -> str:
    vs = [str(v).strip() for v in values if str(v).strip() != ""][:5000]
    if not vs:
        return "text"
    if all(v.lower() in ("y", "n", "true", "false", "예", "아니오", "o", "x") for v in vs) and len({v.lower() for v in vs}) <= 2:
        return "bool"
    if all(re.fullmatch(r"-?\d+", v) for v in vs) and not any(len(v) > 1 and v.startswith("0") for v in vs):
        return "int"
    if all(re.fullmatch(r"-?[\d,]+(\.\d+)?", v) for v in vs):
        return "number"
    if all(re.fullmatch(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}", v) for v in vs):
        return "date"
    if all(re.fullmatch(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}[ T]\d{1,2}:\d{2}(:\d{2})?", v) for v in vs):
        return "datetime"
    return "long_text" if max(len(v) for v in vs) > 200 else "text"


def _slug(h: str) -> str | None:
    s = re.sub(r"[^a-z0-9]+", "_", h.lower()).strip("_")
    return s if s and NAME.match(s) and s not in RESERVED else None


def _project_path(path: str) -> tuple[str, bool]:
    """path relative to the project (cwd), so plans and logs never carry the user's own folder names.
    returns (path_to_store, inside_project)"""
    ab = os.path.abspath(path)
    try:
        rel = os.path.relpath(ab, os.getcwd())
    except ValueError:                         # another drive on Windows
        return Path(ab).name, False
    if rel.startswith(".."):
        return Path(ab).name, False
    return Path(rel).as_posix(), True


def import_plan(path: str, table: str, label: str | None = None, sheet=None) -> dict:
    if not NAME.match(table) or table in RESERVED:
        raise SystemExit(f"--table '{table}': 영문 소문자·숫자·밑줄, 예약어 아님")
    head, rows = _read_table(path, sheet)
    cols = []
    for i, h in enumerate(head):
        vals = [r[i] if i < len(r) else "" for r in rows]
        ty = _infer(vals)
        if ty == "number" and re.search(r"금액|가격|원$|비용|매출|회비|요금|급여|price|amount|fee", h, re.I):
            ty = "money"
        pk = None
        if not NOT_PI.search(h):
            pk = next((k for k, rx in PI_HINT if re.search(rx, h, re.I)), None)
        filled = sum(1 for v in vals if str(v).strip())
        cols.append({"src": h, "name": _slug(h), "label": h, "type": ty,
                     "required": bool(rows) and filled == len(rows), "personal": pk,
                     "samples": ["(개인정보 -- 가림)"] if pk else [str(v) for v in vals[:3]]})
    stored, inside = _project_path(path)
    notes = []
    if not inside:
        notes.append("원본 파일이 프로젝트 밖에 있어 경로 대신 파일 이름만 적었습니다 -- apply 때 --file 로 위치를 알려 주세요")
    return {"format": "data-light-import/1", "file": stored, "sheet": sheet, "table": table,
            "label": label or Path(path).stem, "rows": len(rows), "columns": cols, "approved": False,
            "notes": notes + [
                      "이 계획 파일에는 원본의 실제 값 일부가 들어 있습니다 -- 저장소에 올리지 않습니다 (.gitignore: data/import_plan*.json)",
                      "name 이 null 인 칸은 영문 이름을 정해 채웁니다 (영문 소문자·숫자·밑줄)",
                      "personal 은 개인정보 종류 -- 아니면 null. 개인정보 칸은 pi_ 테이블로 따로 저장됩니다",
                      "사용자 확인 후 approved 를 true 로 바꾸고 import apply 에 사용자 답을 그대로 --approval 로 넘깁니다"]}


def _conv(ty: str, v):
    s = str(v).strip() if v is not None else ""
    if s == "":
        return None
    if ty in ("int",):
        return int(s.replace(",", ""))
    if ty in ("number", "money"):
        return float(s.replace(",", ""))
    if ty == "bool":
        return 1 if s.lower() in ("y", "true", "예", "o", "1") else 0
    if ty == "date":
        y, m, d = re.split(r"[-./]", s)[:3]
        return f"{int(y):04d}-{int(m):02d}-{int(d):02d}"
    return s


def import_apply(plan: dict, model_path: str, db_path: str, approval: str, append: bool = False,
                 file: str | None = None, force_sqlite: bool = False) -> dict:
    if not plan.get("approved"):
        raise SystemExit("approved 가 false 입니다 -- 사용자에게 계획을 보여 주고 확인받은 뒤 true 로 바꾸세요")
    if not (approval or "").strip():
        raise SystemExit("--approval 에 사용자 답을 그대로 넣어 주세요 (확인 기록으로 남습니다)")
    miss = [c["src"] for c in plan["columns"] if not c.get("skip") and not c.get("name")]
    if miss:
        raise SystemExit(f"영문 이름이 없는 칸: {miss}")
    model = load(model_path) if os.path.exists(model_path) else {"format": FORMAT, "project": Path(model_path).parent.name, "db": "sqlite", "tables": []}
    if model.get("db", "sqlite") != "sqlite" and not force_sqlite:
        raise SystemExit("이 정본은 PostgreSQL(Supabase) 용입니다 -- 실제 값이 내 컴퓨터의 SQLite 파일로 남지 않도록 멈춥니다.\n"
                         "Supabase 에는 Table Editor 의 CSV 가져오기로 넣습니다 (개인정보 칸은 pi_ 표로 나눠서). "
                         "그래도 SQLite 에 넣으려면 --force-sqlite")
    src = file or plan["file"]
    if not os.path.exists(src):
        raise SystemExit(f"원본 파일을 찾지 못했습니다: {src} -- 프로젝트 맨 위 폴더에서 실행하거나 --file 로 알려 주세요")
    tn = plan["table"]
    if os.path.exists(db_path) and not append:             # stop before any backup / model change
        con = sqlite3.connect(f"file:{Path(db_path).as_posix()}?mode=ro", uri=True)
        try:
            if con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (tn,)).fetchone():
                have = con.execute(f'SELECT count(*) FROM "{tn}"').fetchone()[0]
                if have:
                    raise SystemExit(f"{tn} 에 이미 {have}행이 있습니다 -- 같은 파일을 두 번 넣지 않도록 멈춥니다. 더하려면 --append")
        finally:
            con.close()
    use = [c for c in plan["columns"] if not c.get("skip")]
    if tn not in tables(model):
        t = {"name": tn, "label": plan.get("label") or tn, "about": f"{Path(plan['file']).name} 에서 가져옴",
             "columns": [{"name": f"{tn}_id", "label": "번호", "type": "id"}]}
        for c in use:
            col = {"name": c["name"], "label": c.get("label") or c["src"], "type": c["type"]}
            if c.get("required"):
                col["required"] = True
            if c.get("personal"):
                col["personal"] = c["personal"]
            t["columns"].append(col)
        model["tables"].append(t)
        split(model)
        errs, _ = check(model)
        if errs:
            raise SystemExit("가져올 테이블 설계에 문제가 있습니다:\n- " + "\n- ".join(errs))
        save(model, model_path)
        created = True
    else:
        created = False
    tb = tables(model)
    main, pi = tb[tn], tb.get("pi_" + tn)
    main_cols = {c["name"] for c in main["columns"]}
    pi_cols = {c["name"] for c in pi["columns"]} if pi else set()
    missing = [c["name"] for c in use if c["name"] not in main_cols | pi_cols]
    if missing:
        raise SystemExit(f"{tn} 에 없는 칸: {missing} -- 정본에 칸을 더하는 것은 설계를 고친 뒤에")
    backup = None
    if os.path.exists(db_path) and os.path.getsize(db_path):
        bdir = Path(db_path).parent / "backup"
        bdir.mkdir(exist_ok=True)
        backup = bdir / f"{Path(db_path).stem}-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}.db"
        bsrc, bdst = sqlite3.connect(db_path), sqlite3.connect(backup)
        try:
            bsrc.backup(bdst)
        finally:
            bsrc.close()
            bdst.close()
    con = sqlite3.connect(db_path)
    con.execute("PRAGMA foreign_keys = ON")
    try:
        exists = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        todo = {t["name"] for t in model["tables"]} - exists
        if todo:
            con.executescript(ddl(model, "sqlite", only=todo))
        head, rows = _read_table(src, plan.get("sheet"))
        idx = {h: i for i, h in enumerate(head)}
        idc = id_column(model, tn)
        bad, n = 0, 0
        with con:
            for r in rows:
                mrow, prow = {}, {}
                for c in use:
                    v = r[idx[c["src"]]] if idx.get(c["src"], 10 ** 9) < len(r) else None
                    try:
                        v = _conv(c["type"], v)
                    except (ValueError, TypeError):
                        bad += 1
                        v = None
                    (prow if c["name"] in pi_cols else mrow)[c["name"]] = v
                cur = con.execute(f"INSERT INTO {tn} ({', '.join(mrow) or idc}) VALUES ({', '.join('?' * len(mrow)) or 'NULL'})",
                                  list(mrow.values()))
                if pi and prow:
                    prow[idc] = cur.lastrowid
                    con.execute(f"INSERT INTO pi_{tn} ({', '.join(prow)}) VALUES ({', '.join('?' * len(prow))})", list(prow.values()))
                n += 1
    finally:
        con.close()
    log = {"at": dt.datetime.now().isoformat(timespec="seconds"), "file": _project_path(file or plan["file"])[0],
           "table": tn, "rows": n,
           "bad_values": bad, "created_table": created, "personal_to": f"pi_{tn}" if pi else None, "backup": _project_path(str(backup))[0] if backup else None,
           "approval": approval.strip()}
    with open(Path(model_path).parent / "import_log.jsonl", "a", encoding="utf-8") as f:
        f.write(json.dumps(log, ensure_ascii=False) + "\n")
    return log


# ----------------------------------------------------------------- import schema (existing SQLite DB -> draft model)
def _light_type(decl: str) -> str:
    d = (decl or "").upper()
    if "BOOL" in d:
        return "bool"
    if "DATETIME" in d or "TIMESTAMP" in d:
        return "datetime"
    if "DATE" in d:
        return "date"
    if "INT" in d:
        return "int"
    if any(k in d for k in ("REAL", "FLOA", "DOUB", "NUMERIC", "DECIMAL")):
        return "number"
    return "text"


def schema_to_model(db_path) -> tuple[dict, list[str]]:
    """read tables / keys / references / NOT NULL / UNIQUE from SQLite. labels stay empty for a person to fill"""
    con = sqlite3.connect(f"file:{Path(db_path).as_posix()}?mode=ro", uri=True)
    issues, out = [], []
    od = {"CASCADE": "delete", "SET NULL": "clear"}
    try:
        names = [r[0] for r in con.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        for tn in names:
            create = con.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name=?", (tn,)).fetchone()[0] or ""
            info = con.execute(f'PRAGMA table_info("{tn}")').fetchall()      # cid, name, type, notnull, dflt, pk
            fks = {r[3]: r for r in con.execute(f'PRAGMA foreign_key_list("{tn}")')}  # from -> (id,seq,table,from,to,upd,del,match)
            single_uq = set()
            combos = []
            for ix in con.execute(f'PRAGMA index_list("{tn}")'):            # seq, name, unique, origin, partial
                if ix[2] and ix[3] in ("u", "c"):
                    cols = [r[2] for r in con.execute(f'PRAGMA index_info("{ix[1]}")')]
                    (single_uq.update(cols) if len(cols) == 1 else combos.append(cols))
            pks = [r for r in info if r[5]]
            cols = []
            names_here = {r[1] for r in info}
            for cid, cn, decl, notnull, dflt, pk in info:
                if cn in ("created_at", "updated_at"):
                    continue
                col = {"name": cn, "label": ""}
                if pk and len(pks) == 1 and "INT" in (decl or "").upper() and cn not in fks:
                    col["type"] = "id"
                elif cn in fks:
                    f = fks[cn]
                    col.update({"type": "ref", "ref": f[2]})
                    if od.get(f[6]):
                        col["on_delete"] = od[f[6]]
                    if pk:
                        col["key"] = True
                    elif notnull:
                        col["required"] = True
                else:
                    col["type"] = _light_type(decl)
                    # SQLite keeps no date/bool/money types -- recover them from CHECK rules and names
                    q = rf'"?{re.escape(cn)}"?'
                    m_in = re.search(rf"CHECK\s*\(\s*{q}\s+IN\s*\(([^)]*)\)\s*\)", create, re.I)
                    if m_in:
                        vals = [v.strip().strip("'\"") for v in m_in.group(1).split(",")]
                        if col["type"] == "int" and sorted(vals) == ["0", "1"]:
                            col["type"] = "bool"
                        else:
                            col["choices"] = [int(v) if col["type"] == "int" and re.fullmatch(r"-?\d+", v) else v for v in vals]
                            if col["type"] not in ("int",):
                                col["type"] = "text"
                    elif re.search(rf"CHECK\s*\(\s*{q}\s*>=\s*0\s*\)", create, re.I) and col["type"] == "number":
                        col["type"] = "money"
                    elif col["type"] == "text" and re.search(r"(^|_)date$|_on$|_day$", cn):
                        col["type"] = "date"
                    elif col["type"] == "text" and re.search(r"_at$|_time$", cn):
                        col["type"] = "datetime"
                    if notnull or pk:
                        col["required"] = True
                    if pk and len(pks) > 1:
                        issues.append(f"{tn}.{cn}: 여러 칸으로 된 키는 다른 표를 가리키는 칸(ref)만 지원 -- 설계를 사람이 확인")
                if cn in single_uq and col["type"] != "id":
                    col["unique"] = True
                if dflt is not None and col["type"] not in ("id", "ref"):
                    v = dflt.strip("'\"")
                    if re.fullmatch(r"-?\d+(\.\d+)?", v):
                        v = float(v) if "." in v else int(v)
                    if col["type"] == "bool" and v in (0, 1):
                        v = bool(v)
                    if not (isinstance(v, str) and v.upper().startswith("CURRENT_")):
                        col["default"] = v
                cols.append(col)
            t = {"name": tn, "label": "", "columns": cols}
            if not ({"created_at", "updated_at"} <= names_here):
                t["timestamps"] = False
            if combos:
                t["unique"] = combos
            if not pks:
                issues.append(f"{tn}: 행을 구별할 키(PRIMARY KEY)가 없습니다 -- {tn}_id 를 더할지 사람이 정함")
            out.append(t)
    finally:
        con.close()
    for t in out:
        for c in t["columns"]:
            if c["type"] in ("text", "long_text", "date") and not NOT_PI.search(c["name"]):
                hints = SCHEMA_PI + (PI_HINT if re.search(r"[^\x00-\x7f]", c["name"]) else [])   # Korean hints only for Korean names
                k = next((k for k, rx in hints if re.search(rx, c["name"], re.I)), None)
                if t["name"].startswith("pi_"):          # already a personal-info table: mark every value column
                    c["personal"] = k or True
                elif k:
                    issues.append(f"{t['name']}.{c['name']}: 개인정보({k})로 보입니다 -- 맞으면 \"personal\": \"{k}\" 표시 후 split")
        bad = [t["name"]] if (not NAME.match(t["name"]) or t["name"] in RESERVED) else []
        bad += [f"{t['name']}.{c['name']}" for c in t["columns"] if not NAME.match(c["name"]) or c["name"] in RESERVED]
        for b in bad:
            issues.append(f"{b}: 이름이 규칙(영어 소문자·숫자·밑줄, 예약어 아님)에 맞지 않습니다 -- 바꾸려면 DB 도 함께 바꿔야 하니 사람이 결정")
    model = {"format": FORMAT, "project": Path(db_path).stem, "db": "sqlite",
             "goals": [], "not_doing": [], "tables": out}
    return model, issues


# ----------------------------------------------------------------- upgrade (to the full framework)
CLASS_TOKENS = {"id", "no", "nm", "cd", "yn", "at", "dt", "amt", "cnt", "val"}


def upgrade(model: dict, schema: str = "app", pi_schema: str | None = None) -> dict:
    """light model → data-expert-framework model.json (passes its check_dictionary / export_all)"""
    errs, _ = check(model)
    if errs:
        raise SystemExit("먼저 check 오류를 고치세요:\n- " + "\n- ".join(errs))
    ty_dom = {"id": ("식별자", "bigint", None), "ref": ("식별자", "bigint", None), "text": ("문자", "text", None),
              "long_text": ("긴글", "text", None), "int": ("정수", "integer", None), "number": ("숫자", "numeric", None),
              "money": ("금액", "numeric(15,2)", ">= 0"), "bool": ("여부", "boolean", None), "date": ("일자", "date", None),
              "datetime": ("시각", "timestamptz", None), "owner": ("사용자ID", "uuid", None)}
    domains: dict[str, dict] = {}

    def dom_of(t, c):
        if c.get("choices"):
            nm = f"선택:{t['name']}.{c['name']}"[:50]
            base = "text" if c["type"] == "text" else "integer"
            domains.setdefault(nm, {"name": nm, "type": base, "check": "IN (" + ", ".join(_q(x) for x in c["choices"]) + ")",
                                    "example": str(c["choices"][0]), "desc": f"{c.get('label', c['name'])} 선택값"})
            return nm
        nm, typ, chk = ty_dom[c["type"]]
        domains.setdefault(nm, {"name": nm, "type": typ, "check": chk, "example": "", "desc": TYPE_KO[c["type"]]})
        return nm

    tbl_out, rels, terms, kor_used = [], [], {}, set()
    tokens: dict[str, None] = {}

    def term(phys, kor, domain, desc):
        if phys in terms:
            return
        k = kor or phys
        if k in kor_used:
            k = f"{k}({phys})"
        kor_used.add(k)
        terms[phys] = {"kor": k, "phys": phys, "word_path": ">".join(phys.split("_")), "domain": domain, "desc": desc or None}
        for tok in phys.split("_"):
            tokens[tok] = None

    for t in ordered(model):
        keys = key_columns(t)
        refs = [c for c in t["columns"] if c["type"] == "ref"]
        kind = "PI" if t["name"].startswith("pi_") else ("거래" if len(refs) >= 2 else "기준")
        cols = []
        for c in all_columns(t):
            d = dom_of(t, c)
            col = {"name": c["name"], "kor": c.get("label") or c["name"], "domain": d,
                   "nullable": not (c.get("required") or c.get("key") or c["type"] in ("id", "owner") or c.get("auto"))}
            if c["name"] in {k["name"] for k in keys}:
                col["pk"] = True
            if c["type"] == "id":
                col["identity"] = True
            if c.get("unique"):
                col["unique"] = True
            if c.get("auto"):
                col["default"] = "now()"
            elif "default" in c:
                col["default"] = _q(c["default"])
            if c["type"] == "ref":
                col["fk"] = {"table": c["ref"], "column": id_column(model, c["ref"]), "on_delete": ON_DELETE[c.get("on_delete", "keep")]}
                rels.append({"from": t["name"], "from_col": c["name"], "to": c["ref"], "to_col": col["fk"]["column"],
                             "type": "1:1" if c.get("key") and len(keys) == 1 else "N:1", "optional": col["nullable"],
                             "on_delete": col["fk"]["on_delete"], "meaning": c.get("label") or c["name"]})
            if c.get("personal"):
                col["pi"] = True
                col["pi_type"] = c["personal"] if isinstance(c["personal"], str) else "개인정보"
            if c.get("about"):
                col["description"] = c["about"]
            cols.append(col)
            term(c["name"], c.get("label"), d, c.get("about"))
        tt = {"name": t["name"], "kor": t.get("label") or t["name"], "kind": kind, "is_pi": t["name"].startswith("pi_"),
              "description": t.get("about") or (t.get("label") or t["name"]), "columns": cols}
        if t.get("unique"):
            tt["uniques"] = [list(u) for u in t["unique"]]
        tbl_out.append(tt)
        term(t["name"], t.get("label"), None, t.get("about"))
    words = [{"kor": tok, "abbr": tok, "full": tok, "is_class": tok in CLASS_TOKENS,
              "desc": "data-light 에서 자동 생성 -- 한글 단어·약어는 공통 사전 대조 때 정한다"} for tok in tokens]
    full = {"project": model.get("project", "data-light project"), "version": "1.0",
            "generated_at": dt.date.today().isoformat(), "db": "postgresql", "db_version": "16",
            "schema": schema, "meta_schema": "meta", "modeling_tool": "pgModeler 1.1",
            "assumptions": [f"data-expert-framework-light ({FORMAT}) 에서 upgrade 로 변환"],
            "open_questions": ["단어(words)의 한글명·약어를 공통 사전(bomiyang-glossary)과 대조해 정리"],
            "words": words, "domains": list(domains.values()), "terms": list(terms.values()),
            "tables": tbl_out, "relationships": rels,
            "pi_suggestions": ["개인정보 칸은 pi_ 테이블에 있음 -- 전문판에서 pi 스키마·암호화(crypto) 적용 검토"],
            "dq_notes": ["light 검사(키·참조·필수·선택값·금액 음수 금지) 통과 상태에서 변환"]}
    if model.get("goals"):
        full["success_goals"] = model["goals"]
        full["dq_notes"].append("성공 기준(success_goals) " + str(len(model["goals"])) + "개 -- 전문판 스모크 테스트로 옮겨 실제 Postgres 에서도 확인")
    if any(c.get("type") == "owner" for t in model["tables"] for c in t["columns"]):
        full["assumptions"].append("owner 칸(로그인한 사람, uuid): Supabase 에서는 DEFAULT auth.uid() -- 전문판의 인증 방식에 맞춰 기본값을 정함")
    if any(t.get("access") for t in model["tables"]):
        full["assumptions"].append("표의 access(members·own·read) 는 Supabase RLS 규칙 -- 전문판의 등급×분류 권한으로 옮길 때 다시 설계: "
                                   + ", ".join(f"{t['name']}={t['access']}" for t in model["tables"] if t.get("access")))
    if model.get("not_doing"):
        full["assumptions"].append("이번 범위 밖(만들지 않음): " + ", ".join(map(str, model["not_doing"])))
    if pi_schema:
        full["pi_schema"] = pi_schema
    return full


# ----------------------------------------------------------------- cli
def main(argv=None) -> int:
    _utf8_stdio()
    ap = argparse.ArgumentParser(prog="datalight", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--version", action="version", version=f"data-light {__version__}")
    sp = ap.add_subparsers(dest="cmd", required=True)
    p = sp.add_parser("check"); p.add_argument("model")
    p = sp.add_parser("split"); p.add_argument("model"); p.add_argument("--write", action="store_true")
    p = sp.add_parser("build"); p.add_argument("model"); p.add_argument("--out"); p.add_argument("--db", choices=["sqlite", "postgresql"])
    p = sp.add_parser("try"); p.add_argument("model")
    p = sp.add_parser("import"); isp = p.add_subparsers(dest="sub", required=True)
    q = isp.add_parser("plan"); q.add_argument("file"); q.add_argument("--table", required=True); q.add_argument("--label")
    q.add_argument("--sheet"); q.add_argument("--out", default="data/import_plan.json")
    q = isp.add_parser("apply"); q.add_argument("plan"); q.add_argument("--model", required=True); q.add_argument("--db", required=True)
    q.add_argument("--approval", required=True); q.add_argument("--append", action="store_true")
    q.add_argument("--file", help="원본 파일 위치 (계획에 파일 이름만 적혔을 때)")
    q.add_argument("--force-sqlite", action="store_true", help="정본이 PostgreSQL 이어도 SQLite 에 넣음")
    q = isp.add_parser("schema"); q.add_argument("db"); q.add_argument("--out", default="data/model.light.json")
    q.add_argument("--force", action="store_true")
    p = sp.add_parser("verify"); p.add_argument("model"); p.add_argument("--db"); p.add_argument("--no-db", action="store_true")
    p.add_argument("--mutate", action="store_true", help="일부러 규칙을 하나씩 빼 보고 검사가 잡는지 확인")
    p = sp.add_parser("hook"); hsp = p.add_subparsers(dest="sub", required=True)
    q = hsp.add_parser("install"); q.add_argument("--repo", default=".")
    g = q.add_mutually_exclusive_group(); g.add_argument("--no-ci", dest="ci", action="store_false", default=None)
    g.add_argument("--ci", dest="ci", action="store_true")
    q.add_argument("--append", action="store_true", help="이미 있는 pre-commit 끝에 덧붙임")
    q = hsp.add_parser("status"); q.add_argument("--repo", default=".")
    p = sp.add_parser("upgrade"); p.add_argument("model"); p.add_argument("--out", default="model.json")
    p.add_argument("--schema", default="app"); p.add_argument("--pi-schema")
    a = ap.parse_args(argv)

    if a.cmd == "check":
        errs, warns = check(load(a.model))
        for w in warns:
            print("참고  " + w)
        for e in errs:
            print("고칠것 " + e)
        print(f"결과: 고칠 것 {len(errs)}개 · 참고 {len(warns)}개" + (" -- 좋아요!" if not errs else ""))
        return 1 if errs else 0
    if a.cmd == "split":
        m = load(a.model)
        moved = split(m)
        print("\n".join(moved) if moved else "나눌 개인정보 칸이 없습니다")
        if a.write and moved:
            save(m, a.model)
            print(f"저장: {a.model}")
        elif moved:
            print("(미리보기 -- 저장하려면 --write)")
        return 0
    if a.cmd == "build":
        m = load(a.model)
        errs, _ = check(m)
        if errs:
            print("check 오류를 먼저 고치세요:\n- " + "\n- ".join(errs))
            return 1
        out = a.out or str(Path(a.model).parent / "out")
        files = build(m, out, a.db)
        print(f"만듦 → {out}: {', '.join(files)}")
        return 0
    if a.cmd == "try":
        m = load(a.model)
        errs, _ = check(m)
        if errs:
            print("check 오류를 먼저 고치세요:\n- " + "\n- ".join(errs))
            return 1
        ok, lines = try_sqlite(m)
        print("\n".join(lines))
        print("결과: " + ("규칙이 모두 제대로 막습니다" if ok else "실패한 줄을 확인하세요"))
        return 0 if ok else 1
    if a.cmd == "import":
        if a.sub == "plan":
            plan = import_plan(a.file, a.table, a.label, a.sheet)
            save(plan, a.out)
            print(f"계획 → {a.out}: {plan['rows']}행, 칸 {len(plan['columns'])}개")
            for c in plan["columns"]:
                print(f"  {c['src']:<16} → {c['name'] or '(이름 필요)':<16} {TYPE_KO[c['type']]:<8}"
                      + (f" 개인정보:{c['personal']}" if c['personal'] else ""))
            for n in plan["notes"][:2]:
                print("참고: " + n)
            return 0
        if a.sub == "schema":
            if os.path.exists(a.out) and not a.force:
                print(f"{a.out} 가 이미 있습니다 -- 덮어쓰려면 --force (다른 이름은 --out)")
                return 1
            m, issues = schema_to_model(a.db)
            save(m, a.out)
            print(f"초안 → {a.out}: 표 {len(m['tables'])}개. 한글 이름표(label)·성공 기준(goals)은 사람과 함께 채웁니다")
            for i in issues:
                print("확인: " + i)
            return 0
        log = import_apply(load(a.plan), a.model, a.db, a.approval, a.append, a.file, a.force_sqlite)
        print(json.dumps(log, ensure_ascii=False, indent=2))
        return 0
    if a.cmd == "verify":
        ok, lines = verify(a.model, a.db, not a.no_db, a.mutate)
        print("\n".join(lines))
        return 0 if ok else 1
    if a.cmd == "hook" and a.sub == "status":
        same, lines = hook_status(a.repo)
        print("\n".join(lines))
        return 0 if same else 1
    if a.cmd == "hook":
        for d in hook_install(a.repo, a.ci, a.append):
            print("만듦: " + d)
        print("이제 정본을 고쳐 커밋하면 verify 가 자동으로 돕니다. 시험: 정본에 일부러 오류를 넣고 커밋해 보세요")
        return 0
    if a.cmd == "upgrade":
        full = upgrade(load(a.model), a.schema, a.pi_schema)
        save(full, a.out)
        print(f"전문판 정본 → {a.out}: 테이블 {len(full['tables'])}개 · 단어 {len(full['words'])}개 "
              "(단어의 한글명·약어는 전문판 공통 사전 대조에서 정리)")
        return 0
    return 2


if __name__ == "__main__":
    sys.exit(main())
