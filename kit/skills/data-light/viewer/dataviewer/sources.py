"""원본 파일 → 보기 모델. 원본마다 형식이 달라도 같은 모양으로 바꾼다.

보기 모델 (원본 하나):
  {"id", "name", "kind": model|xlsx|dbm, "path", "project", "error": None,
   "tables": [{"key": "schema.table", "schema", "name", "kor", "kind", "is_pi", "comment", "description", "meaning",
               "columns": [{"name", "kor", "type", "pk", "fk", "unique", "nullable", "pi", "default", "description", "meaning"}],
               "pos": {"x", "y"} | None}],
   "relations": [{"id", "from", "from_cols", "to", "to_cols", "optional", "one_to_one", "label", "rel_kind"}]}
파일에 적힌 것만 옮긴다. 관계는 원본의 FK(제약)에서만 만든다.
rel_kind (관계선 종류, classify_relations): identifying 식별 · mandatory 비식별 필수 · optional 비식별 선택 · code 코드 참조.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import time
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import unquote, urlsplit

KINDS = ("코드", "기준", "거래", "이력", "PI")


class SourceError(ValueError):
    """원본을 읽을 수 없다 -- 화면에 그대로 보여 줄 문장."""


def source_id(path) -> str:
    return hashlib.sha1(str(Path(path).resolve()).encode("utf-8")).hexdigest()[:10]


def detect(path) -> str:
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".dbm":
        return "dbm"
    if suf in (".xlsx", ".xlsm"):
        return "xlsx"
    if suf == ".json":
        return "model"
    raise SourceError(f"읽을 수 있는 원본이 아니다: {p.name} (model.json · 정의서 .xlsx · pgModeler .dbm)")


_CODE_NAME = re.compile(r"^(code|cd)_|_(code|cd)s?$|^codes?$", re.I)


def _kind(raw, name: str) -> str:
    """원본에 구분이 적혀 있으면 그대로, 없으면(.dbm·실제 DB·light) 이름 관례로: pi_ = 개인정보, *_code·*_cd·code_* = 코드."""
    if raw in KINDS:
        return raw
    if name.startswith("pi_"):
        return "PI"
    return "코드" if _CODE_NAME.search(name) else "기준"


def _rel_id(frm: str, cols, to: str) -> str:
    return f"{frm}({','.join(cols)})->{to}"


# ---------------------------------------------------------------- 의미 (data-modeling scripts/describe.py 와 같은 표기)
# 뷰어는 따로 실려 나가므로(light 사본) 렌더러를 불러오지 않고 같은 규칙을 여기 둔다 -- 바꾸면 둘 다 (tests/test_describe.py 가 맞춰 본다).
MEANING_LABELS = {"grain": "한 행", "scope": "범위", "source_kind": "구분", "history": "변경·이력", "split_reason": "분리 이유",
                  "tech_reason": "도입 이유", "unit": "단위", "value_rule": "값 규칙", "null_meaning": "비어 있으면", "time_zone": "시간대",
                  "code_set": "코드 집합", "key_kind": "식별", "source": "원천", "derivation": "계산·변환", "protection": "보호"}
_KIND_TXT = {"source_kind": {"source": "원천", "derived": "파생", "managed": "업무 입력·관리", "code": "코드", "technical": "기술 관리"},
             "key_kind": {"natural": "업무 자연키", "surrogate": "내부 대리키", "external": "외부 시스템 키"}}
_TAGS = re.compile(r"(\s\[(개인정보|암호화:[^\]]*)\])+$")


def _source_txt(src) -> str:
    if isinstance(src, list):
        return "; ".join(t for t in (_source_txt(x) for x in src) if t)
    if not isinstance(src, dict):
        return "" if src is None else str(src)
    cols = src.get("column")
    cols = cols if isinstance(cols, list) else ([cols] if cols else [])
    out = " ".join(p for p in ((f"자산 {src['asset']}" if src.get("asset") else ""), ("열 " + ", ".join(map(str, cols)) if cols else "")) if p)
    if src.get("rule"):
        out = (out + " → " if out else "") + str(src["rule"])
    if src.get("on_error"):
        out += f" (오류·미상: {src['on_error']})"
    return out


def _meaning(obj: dict) -> dict:
    out = {}
    for k in MEANING_LABELS:
        v = obj.get(k)
        if v in (None, "", []):
            continue
        out[k] = _source_txt(v) if k == "source" else _KIND_TXT.get(k, {}).get(v, v if isinstance(v, str) else json.dumps(v, ensure_ascii=False))
    return out


def parse_comment(text) -> dict:
    """DB COMMENT·.dbm comment ('<한글명> — <설명> [꼬리표]' + '<이름표>: <값>' 줄) -> {kor, description, meaning}."""
    lines = str(text or "").replace("\r\n", "\n").strip().split("\n")
    head, rest = lines[0], lines[1:]
    parts = re.split(r"\s+(?:—|--)\s+", head, maxsplit=1)
    kor = re.sub(r"\s*\[[^\]]*\]", "", parts[0]).strip()
    desc = _TAGS.sub("", parts[1]) if len(parts) > 1 else ""
    by_label = {v: k for k, v in MEANING_LABELS.items()}
    mean, extra = {}, []
    for ln in rest:
        label, sep, v = ln.partition(": ")
        if sep and label in by_label:
            mean[by_label[label]] = v
        elif mean:
            last = list(mean)[-1]
            mean[last] += "\n" + ln
        else:
            extra.append(ln)
    return {"kor": kor, "description": "\n".join([desc] + extra) if extra else desc, "meaning": mean}


# ---------------------------------------------------------------- model.json (정본)
def read_model(path) -> dict:
    try:
        m = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise SourceError(f"정본(model.json)을 읽을 수 없다: {exc}") from None
    if not isinstance(m, dict) or not isinstance(m.get("tables"), list):
        raise SourceError("정본 형식이 아니다 (tables 목록이 없다)")
    if str(m.get("format") or "").startswith("data-light/"):
        return read_light(m, Path(path))
    doms = {d.get("name"): d for d in m.get("domains") or [] if isinstance(d, dict)}
    default = m.get("schema") or "public"

    def schema_of(t):
        if t.get("schema"):
            return t["schema"]
        if m.get("pi_schema") and str(t.get("name", "")).startswith("pi_"):
            return m["pi_schema"]
        return default
    keys = {t["name"]: f"{schema_of(t)}.{t['name']}" for t in m["tables"] if isinstance(t, dict) and t.get("name")}
    tables, rels = [], []
    for t in m["tables"]:
        if not isinstance(t, dict) or not t.get("name"):
            continue
        key = keys[t["name"]]
        cols = []
        for c in t.get("columns") or []:
            if not isinstance(c, dict) or not c.get("name"):
                continue
            fk = c.get("fk") if isinstance(c.get("fk"), dict) else None
            cols.append({"name": c["name"], "kor": c.get("kor") or "", "type": (doms.get(c.get("domain")) or {}).get("type") or c.get("domain") or "",
                         "pk": bool(c.get("pk")), "fk": f"{keys.get(fk.get('table'), fk.get('table'))}.{fk.get('column')}" if fk else None,
                         "unique": bool(c.get("unique")), "nullable": c.get("nullable", True) is not False, "pi": bool(c.get("pi")),
                         "default": "IDENTITY" if c.get("identity") else (None if c.get("default") is None else str(c.get("default"))),
                         "description": c.get("description") or "", "meaning": _meaning(c)})
            if fk and fk.get("table") in keys:
                rels.append({"from": key, "from_cols": [c["name"]], "to": keys[fk["table"]], "to_cols": [fk.get("column")],
                             "optional": c.get("nullable", True) is not False, "one_to_one": bool(c.get("pk")) and sum(1 for x in t.get("columns") or [] if x.get("pk")) == 1,
                             "label": c.get("kor") or c["name"]})
        tables.append({"key": key, "schema": key.split(".")[0], "name": t["name"], "kor": t.get("kor") or "", "kind": _kind(t.get("kind"), t["name"]),
                       "is_pi": bool(t.get("is_pi")) or t["name"].startswith("pi_"), "comment": t.get("description") or "",
                       "description": t.get("description") or "", "meaning": _meaning(t), "columns": cols, "pos": None})
    logical_of = {keys[t["name"]]: [x for x in (t.get("logical") or []) if isinstance(x, str)]
                  for t in m["tables"] if isinstance(t, dict) and t.get("name") in keys}
    for t in tables:
        t["logical"] = logical_of.get(t["key"], [])
        src = next((x for x in m["tables"] if isinstance(x, dict) and keys.get(x.get("name")) == t["key"]), {})
        if src.get("tech_reason"):
            t["tech_reason"] = str(src["tech_reason"])
    return {"project": m.get("project") or Path(path).stem, "tables": tables, "relations": _finish_rels(rels),
            "levels": read_levels(m, logical_of)}


# ---------------------------------------------------------------- 개념·논리 모델 (FW-10)
STATUS_KO = {"draft": "작성 중", "reverse_draft": "역매핑 초안", "confirmed": "확인됨"}


def _h(obj) -> str:
    return hashlib.sha256(json.dumps(obj, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()[:16]


def level_hash(m: dict, level: str) -> str:
    """data-modeling scripts/concept_model.py level_hash 와 같은 지문 (확인 뒤 바뀌었는지 보이려고) -- 바꾸면 둘 다."""
    if level == "concept":
        cm = m.get("concept_model") or {}
        return _h({"scope": cm.get("scope"),
                   "concepts": [{k: c.get(k) for k in ("id", "name", "definition", "kind", "aliases")} for c in cm.get("concepts") or [] if isinstance(c, dict)],
                   "relations": [{k: r.get(k) for k in ("id", "from", "to", "phrase", "cardinality")} for r in cm.get("relations") or [] if isinstance(r, dict)],
                   "terms": [{k: t.get(k) for k in ("original", "source", "concept", "definition", "aliases", "reason", "standard_term")}
                             for t in cm.get("terms") or [] if isinstance(t, dict)]})
    lm = m.get("logical_model") or {}
    return _h({"scope": lm.get("scope"),
               "entities": [{**{k: e.get(k) for k in ("id", "name", "definition", "identifier")},
                             "attributes": [{k: a.get(k) for k in ("id", "name", "definition", "domain", "standard_term")} for a in e.get("attributes") or [] if isinstance(a, dict)],
                             **{k: e.get(k) for k in ("concepts", "concept_reason", "tech_reason", "not_implemented")}}
                            for e in lm.get("entities") or [] if isinstance(e, dict)],
               "relations": [{k: r.get(k) for k in ("id", "from", "to", "phrase", "cardinality", "optional")} for r in lm.get("relations") or [] if isinstance(r, dict)]})


def _pending(cm: dict) -> int:
    """concept_model.py 의 '사용자 검토 대상' 수 (AI 개념 확인 전 · 검토 대기 정리 · 확인 안 된 원문 묶음) -- 있으면 확인됨으로 보이지 않는다."""
    concepts = [c for c in cm.get("concepts") or [] if isinstance(c, dict)]
    terms = [t for t in cm.get("terms") or [] if isinstance(t, dict)]
    n = sum(1 for c in concepts if c.get("origin") == "ai" and c.get("status") != "confirmed")
    n += sum(1 for t in terms if t.get("status") == "검토 대기")
    groups = {}
    for t in terms:
        if t.get("concept"):
            groups.setdefault(t["concept"], []).append(t)
    n += sum(1 for ts in groups.values() if len({t.get("original") for t in ts}) > 1 and any(t.get("status") != "확인" for t in ts))
    return n


def _level_head(m: dict, level: str) -> dict | None:
    blk = m.get(f"{level}_model")
    if not isinstance(blk, dict):
        return None
    st = blk.get("status")
    conf = blk.get("confirmed") if isinstance(blk.get("confirmed"), dict) else {}
    changed = st == "confirmed" and conf.get("hash") != level_hash(m, level)
    if level == "concept" and st == "confirmed" and not changed and _pending(blk):
        return {"status": "draft", "label": f"확인됨이었으나 검토 대상 {_pending(blk)}개가 생겼다 -- 정하고 다시 확인", "version": blk.get("version"),
                "scope": blk.get("scope"), "confirmed_by": None, "confirmed_at": None}
    eff = "draft" if changed else st
    return {"status": eff, "label": STATUS_KO.get(eff, str(eff or "상태 없음")) + (" (확인 뒤 바뀜 -- 다시 확인 필요)" if changed else ""),
            "version": blk.get("version"), "scope": blk.get("scope"), "confirmed_by": conf.get("by") if not changed else None,
            "confirmed_at": conf.get("at") if not changed else None}


def read_levels(m: dict, logical_of: dict) -> dict:
    """{"concept": {...} | None, "logical": {...} | None} -- 없으면 None (화면은 '미작성', 물리로 대신하지 않는다)."""
    out = {"concept": None, "logical": None}
    lm = m.get("logical_model") if isinstance(m.get("logical_model"), dict) else None
    ents = [e for e in (lm or {}).get("entities") or [] if isinstance(e, dict) and e.get("id")]
    tables_of = {}
    for tk, ls in logical_of.items():
        for lid in ls:
            tables_of.setdefault(lid, []).append(tk)
    cm = m.get("concept_model") if isinstance(m.get("concept_model"), dict) else None
    if cm is not None:
        head = _level_head(m, "concept")
        terms = [t for t in cm.get("terms") or [] if isinstance(t, dict)]
        nodes = []
        for c in cm.get("concepts") or []:
            if not isinstance(c, dict) or not c.get("id"):
                continue
            nodes.append({"key": c["id"], "name": str(c.get("name") or ""), "definition": str(c.get("definition") or ""),
                          "kind": str(c.get("kind") or ""), "aliases": [str(a) for a in c.get("aliases") or []],
                          "origin": c.get("origin"), "status": c.get("status"),
                          "terms": [{k: x.get(k) for k in ("original", "source", "reason", "status", "standard_term")}
                                    for x in terms if x.get("concept") == c["id"]],
                          "maps": [e["id"] for e in ents if c["id"] in (e.get("concepts") or [])]})
        ids = {n["key"] for n in nodes}
        rels = [{"id": str(r.get("id")), "from": r["from"], "to": r["to"], "phrase": str(r.get("phrase") or ""),
                 "cardinality": r.get("cardinality") if r.get("cardinality_status") == "confirmed" else None}
                for r in cm.get("relations") or [] if isinstance(r, dict) and r.get("from") in ids and r.get("to") in ids]
        out["concept"] = {**head, "nodes": nodes, "relations": rels,
                          "unmapped_terms": [t.get("original") for t in terms if not t.get("concept")]}
    if lm is not None:
        head = _level_head(m, "logical")
        nodes = [{"key": e["id"], "name": str(e.get("name") or ""), "definition": str(e.get("definition") or ""),
                  "identifier": [str(x) for x in e.get("identifier") or []],
                  "attributes": [{"name": str(a.get("name") or ""), "definition": str(a.get("definition") or "")}
                                 for a in e.get("attributes") or [] if isinstance(a, dict)],
                  "concepts": [x for x in e.get("concepts") or [] if isinstance(x, str)], "tables": sorted(tables_of.get(e["id"], [])),
                  "reason": e.get("concept_reason") or e.get("tech_reason")} for e in ents]
        ids = {n["key"] for n in nodes}
        rels = [{"id": str(r.get("id")), "from": r["from"], "to": r["to"], "phrase": str(r.get("phrase") or ""),
                 "cardinality": r.get("cardinality"), "optional": bool(r.get("optional"))}
                for r in lm.get("relations") or [] if isinstance(r, dict) and r.get("from") in ids and r.get("to") in ids]
        out["logical"] = {**head, "nodes": nodes, "relations": rels}
    return out


LIGHT_TYPES = {"id": "자동 번호", "ref": "가리킴", "text": "짧은 글", "long_text": "긴 글", "int": "정수", "number": "숫자",
               "money": "금액", "bool": "예/아니오", "date": "날짜", "datetime": "날짜+시각", "owner": "만든 사람"}


def read_light(m: dict, path: Path) -> dict:
    """data-expert-framework-light 의 data/model.light.json (data-light/1) -- 표·칸·가리킴(ref)·개인정보(personal)."""
    schema = "public" if m.get("db") == "postgresql" else "main"
    tabs = [t for t in m["tables"] if isinstance(t, dict) and isinstance(t.get("name"), str)]
    key_of = {}
    for t in tabs:
        cols = [c for c in t.get("columns") or [] if isinstance(c, dict)]
        ids = [c.get("name") for c in cols if c.get("type") == "id"]
        keys = [c.get("name") for c in cols if c.get("type") == "ref" and c.get("key")]
        key_of[t["name"]] = ids or keys
    tables, rels = [], []
    for t in tabs:
        key = f"{schema}.{t['name']}"
        cols = []
        pks = key_of[t["name"]]
        for c in t.get("columns") or []:
            if not isinstance(c, dict) or not isinstance(c.get("name"), str):
                continue
            ty = str(c.get("type") or "")
            fk = None
            if ty == "ref" and isinstance(c.get("ref"), str):
                tgt = key_of.get(c["ref"]) or [f"{c['ref']}_id"]
                fk = f"{schema}.{c['ref']}.{tgt[0]}"
                if c["ref"] in key_of:
                    rels.append({"from": key, "from_cols": [c["name"]], "to": f"{schema}.{c['ref']}", "to_cols": [tgt[0]],
                                 "optional": not c.get("required") and not c.get("key"), "one_to_one": pks == [c["name"]],
                                 "label": c.get("label") or c["name"]})
            choices = c.get("choices") if isinstance(c.get("choices"), list) else None
            cols.append({"name": c["name"], "kor": c.get("label") or "",
                         "type": ty + (f"({'·'.join(map(str, choices))})" if choices else ""),
                         "pk": c["name"] in pks, "fk": fk, "unique": bool(c.get("unique")),
                         "nullable": not (c.get("required") or ty in ("id", "owner") or c["name"] in pks), "pi": bool(c.get("personal")),
                         "default": None if c.get("default") is None else str(c.get("default")),
                         "description": str(c.get("about") or ""), "meaning": {}})
        if t.get("timestamps", True) is not False:
            for n, k in (("created_at", "만든 시각"), ("updated_at", "고친 시각")):
                cols.append({"name": n, "kor": k, "type": "datetime", "pk": False, "fk": None, "unique": False,
                             "nullable": False, "pi": False, "default": None, "description": "", "meaning": {}})
        name = t["name"]
        tables.append({"key": key, "schema": schema, "name": name, "kor": t.get("label") or "", "kind": _kind(None, name),
                       "is_pi": name.startswith("pi_"), "comment": t.get("about") or "", "description": t.get("about") or "", "meaning": {},
                       "columns": cols, "pos": None})
    return {"project": m.get("project") or path.stem, "tables": tables, "relations": _finish_rels(rels), "light": True}


# ---------------------------------------------------------------- 정의서 엑셀 (export_all 의 04_table_spec.xlsx)
def _rows(ws):
    rows = [tuple("" if v is None else v for v in r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return [], []
    head = [str(h).strip() for h in rows[0]]
    return head, [dict(zip(head, r)) for r in rows[1:] if any(str(v).strip() for v in r)]


def read_xlsx(path) -> dict:
    try:
        import openpyxl
    except ImportError:
        raise SourceError("정의서 엑셀을 읽으려면 openpyxl 이 필요하다 (pip install openpyxl)") from None
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 -- 열 수 없는 파일은 모두 같은 안내
        raise SourceError(f"엑셀을 열 수 없다: {exc}") from None
    try:
        need = ("테이블목록", "컬럼명세")
        if not all(n in wb.sheetnames for n in need):
            raise SourceError("이 저장소의 정의서 양식이 아니다 -- '테이블목록'·'컬럼명세' 시트가 있어야 한다 (export_all.py 가 만든 04_table_spec.xlsx)")
        _, tlist = _rows(wb["테이블목록"])
        _, clist = _rows(wb["컬럼명세"])
        project = Path(path).stem
        if "개요" in wb.sheetnames:
            for r in _rows(wb["개요"])[1]:
                if str(r.get("항목")) == "프로젝트" and r.get("내용"):
                    project = str(r["내용"])
    finally:
        wb.close()
    keys, tables, dup = {}, {}, set()
    for r in tlist:
        name = str(r.get("테이블명") or "").strip()
        if not name:
            continue
        schema = str(r.get("스키마") or "public").strip()
        if name in keys:
            dup.add(name)
            continue
        keys[name] = f"{schema}.{name}"
        tables[name] = {"key": keys[name], "schema": schema, "name": name, "kor": str(r.get("논리명") or ""), "kind": _kind(r.get("구분"), name),
                        "is_pi": str(r.get("PI여부")).upper() == "Y" or name.startswith("pi_"), "comment": str(r.get("설명") or ""),
                        "description": str(r.get("설명") or ""),
                        "meaning": {k: str(r[h]) for k, h in (("grain", "한 행의 기준"), ("scope", "범위"), ("source_kind", "원천/파생"),
                                                               ("history", "변경·이력")) if r.get(h) not in (None, "")},
                        "columns": [], "pos": None}
    rels, closed = [], set()
    for r in clist:
        tname = str(r.get("테이블명") or "").strip()
        if tname not in tables or not str(r.get("컬럼명") or "").strip() or tname in closed:
            continue
        if tname in dup:                                # 같은 이름이 또 나오면(순번이 1로 돌아가거나 컬럼이 겹치면) 그 뒤는 다른 스키마의 것
            seen = {c["name"] for c in tables[tname]["columns"]}
            if seen and (str(r.get("순번")).strip() == "1" or str(r["컬럼명"]).strip() in seen):
                closed.add(tname)
                continue
        fk = str(r.get("FK(참조)") or "").strip() or None
        col = {"name": str(r["컬럼명"]).strip(), "kor": str(r.get("논리명") or ""), "type": str(r.get("자료형") or ""),
               "pk": str(r.get("PK")).upper() == "PK", "fk": None, "unique": str(r.get("UNIQUE")).upper() in ("UK", "Y"),
               "nullable": str(r.get("NULL")).upper() != "N", "pi": str(r.get("PI")).upper() == "Y",
               "default": None if r.get("기본값") in (None, "") else str(r.get("기본값")), "description": str(r.get("설명") or ""),
               "meaning": {k: str(r[h]) for k, h in (("unit", "단위"), ("null_meaning", "NULL 의미"), ("source", "원천·변환"))
                           if r.get(h) not in (None, "")}}
        if fk and "." in fk:
            ft, fc = fk.rsplit(".", 1)
            col["fk"] = f"{keys.get(ft, ft)}.{fc}"
            if ft in keys:
                rels.append({"from": keys[tname], "from_cols": [col["name"]], "to": keys[ft], "to_cols": [fc], "optional": col["nullable"],
                             "one_to_one": False, "label": col["kor"] or col["name"]})
        tables[tname]["columns"].append(col)
    for rel in rels:                                    # PK 하나가 곧 FK 면 1:1 (pi_ 테이블)
        t = tables[rel["from"].split(".", 1)[1]]
        pks = [c["name"] for c in t["columns"] if c["pk"]]
        rel["one_to_one"] = pks == rel["from_cols"]
    out = {"project": project, "tables": list(tables.values()), "relations": _finish_rels(rels)}
    if dup:
        out["warning"] = f"같은 이름의 테이블이 여러 스키마에 있다 ({', '.join(sorted(dup))}) -- 정의서의 컬럼명세에는 스키마가 없어 처음 것만 보인다. 정본(model.json)을 열면 모두 보인다"
    return out


# ---------------------------------------------------------------- pgModeler .dbm
def _dbm_type(el) -> str:
    if el is None:
        return ""
    name = el.get("name") or ""
    if el.get("with-timezone") == "true" and name in ("timestamp", "time"):
        name += " with time zone"
    ln, pr, dim = el.get("length") or "0", el.get("precision"), int(el.get("dimension") or 0)
    if name.startswith(("timestamp", "time")) and pr:
        head, _, tail = name.partition(" ")
        base = f"{head}({pr})" + (f" {tail}" if tail else "")
    elif ln not in ("", "0"):
        base = f"{name}({ln}" + (f",{pr}" if pr else "") + ")"
    else:
        base = name
    return base + "[]" * dim


def read_dbm(path) -> dict:
    try:
        raw = Path(path).read_bytes()
    except OSError as exc:
        raise SourceError(f".dbm 을 읽을 수 없다: {exc}") from None
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff") or raw[:4] in (b"\x00<\x00?", b"<\x00?\x00"):
        raise SourceError(".dbm 이 UTF-16 이다 -- pgModeler 는 UTF-8 로 쓴다 (읽지 않는다)")
    if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
        raise SourceError(".dbm 에 DOCTYPE·ENTITY 선언이 있다 -- pgModeler 가 만든 파일이 아니다 (읽지 않는다)")
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise SourceError(f".dbm XML 이 깨졌다: {exc}") from None
    if root.tag != "dbmodel":
        raise SourceError(".dbm 형식이 아니다 (맨 위가 dbmodel 이 아니다)")
    tables, rels = {}, []
    fk_items = []

    def comment_of(el):
        c = el.find("comment")
        return (c.text or "").strip() if c is not None else ""
    for t in list(root.findall("table")) + list(root.findall("foreigntable")):
        sch = (t.find("schema").get("name") if t.find("schema") is not None else "public")
        key = f"{sch}.{t.get('name')}"
        pos = t.find("position")
        cols = {}
        for c in t.findall("column"):
            pc = parse_comment(comment_of(c))
            cols[c.get("name")] = {"name": c.get("name"), "kor": c.get("alias") or pc["kor"], "type": _dbm_type(c.find("type")), "pk": False,
                                   "fk": None, "unique": False, "nullable": c.get("not-null") != "true", "pi": "[개인정보]" in comment_of(c),
                                   "default": "IDENTITY" if c.get("identity-type") else c.get("default-value"),
                                   "description": pc["description"], "meaning": pc["meaning"]}
        for k in t.findall("constraint"):
            names = [n for cs in k.findall("columns") if cs.get("ref-type") == "src-columns" for n in (cs.get("names") or "").split(",") if n]
            if k.get("type") == "pk-constr":
                for n in names:
                    if n in cols:
                        cols[n]["pk"] = True
            elif k.get("type") == "uq-constr" and len(names) == 1 and names[0] in cols:
                cols[names[0]]["unique"] = True
            elif k.get("type") == "fk-constr":
                fk_items.append((key, k))
        name = t.get("name")
        tables[key] = {"key": key, "schema": sch, "name": name, "kor": t.get("alias") or "", "kind": _kind(None, name),
                       "is_pi": name.startswith("pi_"), "comment": comment_of(t), "columns": list(cols.values()),
                       "description": parse_comment(comment_of(t))["description"], "meaning": parse_comment(comment_of(t))["meaning"],
                       "pos": {"x": float(pos.get("x") or 0), "y": float(pos.get("y") or 0)} if pos is not None else None}
    fk_items += [(k.get("table"), k) for k in root.findall("constraint") if k.get("type") == "fk-constr"]
    for owner, k in fk_items:
        src = [n for cs in k.findall("columns") if cs.get("ref-type") == "src-columns" for n in (cs.get("names") or "").split(",") if n]
        dst = [n for cs in k.findall("columns") if cs.get("ref-type") == "dst-columns" for n in (cs.get("names") or "").split(",") if n]
        frm, to = owner or k.get("table"), k.get("ref-table")
        if frm not in tables or to not in tables:
            continue
        cols = {c["name"]: c for c in tables[frm]["columns"]}
        for s, d in zip(src, dst):
            if s in cols:
                cols[s]["fk"] = f"{to}.{d}"
        pks = [c["name"] for c in tables[frm]["columns"] if c["pk"]]
        first = cols.get(src[0], {}) if src else {}
        rels.append({"from": frm, "from_cols": src, "to": to, "to_cols": dst,
                     "optional": any(cols.get(s, {}).get("nullable", True) for s in src), "one_to_one": pks == src,
                     "label": first.get("kor") or k.get("alias") or (src[0] if len(src) == 1 else k.get("name") or "")})
    out = {"project": Path(path).stem, "tables": list(tables.values()), "relations": _finish_rels(rels)}
    # pgModeler 에서 직접 그린 관계(1:1·1:n·n:n·상속 …)는 FK 제약이 파일에 없다 -- 짐작해 긋지 않고 알린다
    native = [r.get("name") or "?" for r in root.findall("relationship") if r.get("type") not in (None, "relfk")]
    if native:
        out["warning"] = (f"pgModeler 관계 {len(native)}개는 FK 제약이 파일에 없어 그리지 않았다 ({', '.join(native[:5])}"
                          f"{' …' if len(native) > 5 else ''}) -- 그 관계가 만드는 컬럼도 파일에 없다. pgModeler 에서 '관계를 FK 로 바꾸기' 뒤 저장하거나 정본에서 만든 .dbm 을 연다")
    return out


REL_KINDS = ("identifying", "mandatory", "optional", "code")


def _rel_kind(identifying: bool, code: bool, optional: bool) -> str:
    if identifying:
        return "identifying"
    if code:
        return "code"
    return "optional" if optional else "mandatory"


def classify_relations(view: dict) -> dict:
    """관계선 종류를 붙인다 (화면에서 종류별로 켜고 끄고, 보이는 선만으로 다시 배치하려고).
    물리: FK 컬럼이 모두 자식의 PK 안에 있으면 식별(부모 식별자를 물려받음) -- 자기 참조는 식별이 아니다.
          아니면 부모가 코드 테이블이면 코드 참조, 아니면 FK 컬럼이 비어도 되는지(optional)로 필수/선택.
    논리: 한쪽 엔터티의 식별자가 상대 엔터티(ID 또는 이름)를 품으면 식별. 상대가 코드 테이블에만 대응하면 코드 참조.
    개념: 나누지 않는다 (업무 관계만 있다)."""
    tabs = {t["key"]: t for t in view.get("tables") or []}
    for r in view.get("relations") or []:
        child, parent = tabs.get(r["from"]), tabs.get(r["to"])
        pks = {c["name"] for c in (child or {}).get("columns") or [] if c.get("pk")}
        cols = [c for c in r.get("from_cols") or [] if c]
        ident = bool(cols) and r["from"] != r["to"] and all(c in pks for c in cols)
        r["rel_kind"] = _rel_kind(ident, (parent or {}).get("kind") == "코드", bool(r.get("optional")))
    lg = (view.get("levels") or {}).get("logical")
    if lg:
        nodes = {n["key"]: n for n in lg.get("nodes") or []}

        def holds(a, b):        # a 의 식별자가 b 를 품는가
            na, nb = nodes.get(a) or {}, nodes.get(b) or {}
            return bool(nb) and a != b and bool({nb.get("key"), nb.get("name")} & set(na.get("identifier") or []))

        def is_code(e):
            ts = (nodes.get(e) or {}).get("tables") or []
            return bool(ts) and all((tabs.get(t) or {}).get("kind") == "코드" for t in ts)
        for r in lg.get("relations") or []:
            ident = holds(r["from"], r["to"]) or holds(r["to"], r["from"])
            r["rel_kind"] = _rel_kind(ident, is_code(r["to"]), bool(r.get("optional")))
    return view


def _finish_rels(rels: list[dict]) -> list[dict]:
    out, seen = [], set()
    for r in rels:
        rid = _rel_id(r["from"], r["from_cols"], r["to"])
        if rid in seen:
            continue
        seen.add(rid)
        out.append(dict(r, id=rid))
    return out


# ---------------------------------------------------------------- 실제 DB (PostgreSQL, 읽기만)
DB_SCHEMES = ("postgresql://", "postgres://")
SYSTEM_SCHEMAS = ("pg_catalog", "information_schema", "pg_toast")
_USER_NS = "n.nspname NOT IN ('pg_catalog','information_schema') AND n.nspname NOT LIKE 'pg\\_%'"
_PG_SQL = {
    "tables": f"""SELECT n.nspname, c.relname, obj_description(c.oid, 'pg_class') FROM pg_class c JOIN pg_namespace n ON n.oid = c.relnamespace
                  WHERE c.relkind IN ('r','p','f') AND {_USER_NS} ORDER BY 1, 2""",
    "columns": f"""SELECT n.nspname, c.relname, a.attname, format_type(a.atttypid, a.atttypmod), a.attnotnull,
                   pg_get_expr(d.adbin, d.adrelid), a.attidentity, col_description(c.oid, a.attnum)
                   FROM pg_attribute a JOIN pg_class c ON c.oid = a.attrelid JOIN pg_namespace n ON n.oid = c.relnamespace
                   LEFT JOIN pg_attrdef d ON d.adrelid = a.attrelid AND d.adnum = a.attnum
                   WHERE c.relkind IN ('r','p','f') AND a.attnum > 0 AND NOT a.attisdropped AND {_USER_NS} ORDER BY 1, 2, a.attnum""",
    "constraints": f"""SELECT n.nspname, c.relname, k.conname, k.contype,
                       ARRAY(SELECT a.attname::text FROM unnest(k.conkey) WITH ORDINALITY u(n, o)
                             JOIN pg_attribute a ON a.attrelid = k.conrelid AND a.attnum = u.n ORDER BY u.o),
                       fn.nspname, fc.relname,
                       ARRAY(SELECT a.attname::text FROM unnest(k.confkey) WITH ORDINALITY u(n, o)
                             JOIN pg_attribute a ON a.attrelid = k.confrelid AND a.attnum = u.n ORDER BY u.o)
                       FROM pg_constraint k JOIN pg_class c ON c.oid = k.conrelid JOIN pg_namespace n ON n.oid = c.relnamespace
                       LEFT JOIN pg_class fc ON fc.oid = k.confrelid LEFT JOIN pg_namespace fn ON fn.oid = fc.relnamespace
                       WHERE k.contype IN ('p','u','f') AND {_USER_NS} ORDER BY 1, 2, 3""",
}


def split_dsn(spec: str) -> tuple[str, list | None]:
    """'postgresql://…#스키마,스키마' -> (접속 문자열, 스키마). 비밀번호에 '#' 이 있을 수 있어 마지막 '#' 에서 나누고,
    뒤가 스키마 이름 모양(글자·숫자·밑줄·쉼표)이 아니면 나누지 않는다 (비밀번호 조각을 스키마로 보이지 않게)."""
    head, sep, tail = spec.rpartition("#")
    if sep and re.fullmatch(r"[A-Za-z_][\w$]*(\s*,\s*[A-Za-z_][\w$]*)*", tail.strip()):
        return head, [x.strip() for x in tail.split(",") if x.strip()]
    return spec, None


def _secrets(dsn: str) -> list[str]:
    try:
        pw = urlsplit(dsn).password
    except ValueError:
        pw = None
    raw = re.match(r"^[a-z]+://[^:/@]*:([^@]*)@", dsn)
    out = {x for x in (pw, unquote(pw) if pw else None, raw.group(1) if raw else None) if x}
    return sorted(out, key=len, reverse=True)


def scrub(text: str, dsn: str) -> str:
    """오류 문장에서 비밀번호(그대로·풀어 쓴 것)를 가린다."""
    for s_ in _secrets(dsn):
        text = text.replace(s_, "***")
    return text


def mask_dsn(dsn: str) -> str:
    """화면·기록에 쓸 이름 -- 비밀번호와 접속 옵션(쿼리)은 빼고 사용자@호스트:포트/DB 만."""
    try:
        u = urlsplit(split_dsn(dsn)[0])
        _ = u.port
    except ValueError:
        return "postgresql://(접속 문자열을 읽을 수 없다)"
    host = u.hostname or "localhost"
    port = f":{u.port}" if u.port else ""
    user = f"{unquote(u.username)}@" if u.username else ""
    return f"postgresql://{user}{host}{port}{u.path or '/'}"


def _kor_comment(text) -> str:
    """COMMENT 의 첫 줄에서 [변경모드:...] 같은 꼬리표를 뺀 한글 이름."""
    if not text:
        return ""
    line = str(text).strip().splitlines()[0]
    line = re.split(r"\s+(?:—|--)\s+", line, maxsplit=1)[0]          # export_all: '학생 — 설명'
    return re.sub(r"\s*\[[^\]]*\]", "", line).strip()


def read_pg(dsn: str, schemas: list | None = None) -> dict:
    """실제 PostgreSQL 의 카탈로그(pg_catalog)만 읽는다 -- 데이터 행은 읽지 않고, 읽기 전용 트랜잭션으로."""
    try:
        import psycopg
    except ImportError:
        raise SourceError("실제 DB 를 읽으려면 psycopg 가 필요하다 (pip install 'psycopg[binary]')") from None
    try:
        with psycopg.connect(dsn, connect_timeout=5, autocommit=False, options="-c statement_timeout=10000") as con:
            con.execute("SET TRANSACTION READ ONLY")
            rows = {k: con.execute(q).fetchall() for k, q in _PG_SQL.items()}
            con.rollback()
    except psycopg.Error as exc:
        msg = str(exc).strip().splitlines()[0] if str(exc).strip() else type(exc).__name__
        raise SourceError(f"DB 에 접속하거나 읽을 수 없다 ({mask_dsn(dsn)}): {scrub(msg, dsn)}") from None
    except Exception as exc:  # noqa: BLE001 -- 접속 문자열 해석 오류 등도 비밀번호 없이
        raise SourceError(f"DB 에 접속하거나 읽을 수 없다 ({mask_dsn(dsn)}): {scrub(type(exc).__name__ + ': ' + str(exc), dsn)}") from None
    keep = (lambda sch: sch not in SYSTEM_SCHEMAS and (not schemas or sch in schemas))
    tables = {}
    for sch, name, comment in rows["tables"]:
        if keep(sch):
            pc = parse_comment(comment)
            tables[f"{sch}.{name}"] = {"key": f"{sch}.{name}", "schema": sch, "name": name, "kor": _kor_comment(comment),
                                       "kind": _kind(None, name), "is_pi": name.startswith("pi_"),
                                       "comment": str(comment or ""), "description": pc["description"], "meaning": pc["meaning"],
                                       "columns": [], "pos": None}
    cols = {}
    for sch, tname, col, ty, notnull, default, ident, comment in rows["columns"]:
        key = f"{sch}.{tname}"
        if key not in tables:
            continue
        pc = parse_comment(comment)
        c = {"name": col, "kor": _kor_comment(comment), "type": ty, "pk": False, "fk": None, "unique": False,
             "nullable": not notnull, "pi": "[개인정보]" in str(comment or ""), "default": "IDENTITY" if ident else default,
             "description": pc["description"], "meaning": pc["meaning"]}
        tables[key]["columns"].append(c)
        cols[(key, col)] = c
    rels = []
    for sch, tname, cname, ctype, src, fsch, ftab, dst in rows["constraints"]:
        key = f"{sch}.{tname}"
        if key not in tables:
            continue
        src, dst = list(src or []), list(dst or [])
        if ctype == "p":
            for n in src:
                if (key, n) in cols:
                    cols[(key, n)]["pk"] = True
        elif ctype == "u" and len(src) == 1 and (key, src[0]) in cols:
            cols[(key, src[0])]["unique"] = True
        elif ctype == "f":
            to = f"{fsch}.{ftab}"
            for s_, d_ in zip(src, dst):
                if (key, s_) in cols:
                    cols[(key, s_)]["fk"] = f"{to}.{d_}"
            if to in tables:
                first = cols.get((key, src[0]), {}) if src else {}
                rels.append({"from": key, "from_cols": src, "to": to, "to_cols": dst,
                             "optional": any(cols.get((key, s_), {}).get("nullable", True) for s_ in src),
                             "one_to_one": False, "label": first.get("kor") or (src[0] if len(src) == 1 else cname)})
    for r in rels:
        pks = [c["name"] for c in tables[r["from"]]["columns"] if c["pk"]]
        r["one_to_one"] = pks == r["from_cols"]
    return classify_relations({"project": mask_dsn(dsn).rsplit("/", 1)[-1] or mask_dsn(dsn), "tables": list(tables.values()), "relations": _finish_rels(rels)})


READERS = {"model": read_model, "xlsx": read_xlsx, "dbm": read_dbm}


def is_db(spec) -> bool:
    return isinstance(spec, str) and spec.lower().startswith(DB_SCHEMES)


class FileSource:
    """원본 파일 -- 수정 시각·크기로 바뀜을 알아챈다. 위치 파일은 원본 옆."""

    def __init__(self, path):
        self.path = Path(path).resolve()
        self.kind = detect(self.path)
        self.id = source_id(self.path)
        self.name = self.path.name
        self.display = str(self.path)
        self.layout_base = self.path

    def stamp(self):
        try:
            st = self.path.stat()
            return (st.st_mtime_ns, st.st_size)
        except OSError:
            return None

    def read(self) -> dict:
        return read_source(self.path)

    def missing(self) -> str | None:
        return None if self.stamp() is not None else f"파일이 없다: {self.name} (생기면 다시 읽는다)"

    def refresh_now(self):
        pass


class DbSource:
    """실제 PostgreSQL (읽기만) -- 카탈로그를 db_poll 초마다 읽어 지문이 바뀌면 다시 그린다.
    접속 문자열 끝의 `#스키마,스키마` 로 볼 스키마를 고른다. 비밀번호는 화면·위치 파일에 남기지 않는다."""

    def __init__(self, dsn: str, *, layout_dir=None, db_poll: float = 5.0):
        if not is_db(dsn):
            raise SourceError("DB 원본은 postgresql:// 로 시작한다")
        self.dsn, self.schemas = split_dsn(dsn)
        self.kind = "db"
        self.display = mask_dsn(self.dsn) + (f" ({', '.join(self.schemas)})" if self.schemas else "")
        self.id = hashlib.sha1(self.display.encode("utf-8")).hexdigest()[:10]
        hostdb = mask_dsn(self.dsn).split("@")[-1].replace("postgresql://", "")
        safe = re.sub(r"[^\w.-]+", "_", hostdb).strip("_") or "db"
        self.name = self.display
        self.layout_base = Path(layout_dir or os.getcwd()).resolve() / f"{safe}.db"
        self.db_poll = db_poll
        self._cache = None              # (읽은 시각, 지문, 보기 모델 또는 SourceError)

    def _load(self, force: bool = False):
        now = time.monotonic()
        if force or self._cache is None or now - self._cache[0] >= self.db_poll:
            try:
                view = read_pg(self.dsn, self.schemas)
                fp = hashlib.sha1(json.dumps(view, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()
                self._cache = (now, fp, view)
            except SourceError as exc:
                self._cache = (now, "error:" + str(exc), exc)
        return self._cache

    def stamp(self):
        return self._load()[1]

    def read(self) -> dict:
        got = self._load()[2]
        if isinstance(got, SourceError):
            raise got
        return dict(got, tables=[dict(t) for t in got["tables"]])

    def missing(self) -> str | None:
        return None

    def refresh_now(self):
        self._load(force=True)


def make_source(spec, *, layout_dir=None):
    return DbSource(spec, layout_dir=layout_dir) if is_db(spec) else FileSource(spec)


def read_source(path) -> dict:
    """원본 하나 → 보기 모델 (읽을 수 없으면 SourceError)."""
    kind = detect(path)
    try:
        view = classify_relations(READERS[kind](path))
    except SourceError:
        raise
    except Exception as exc:  # noqa: BLE001 -- 형식이 어긋난 원본(이름이 숫자 · 위치가 글자 …)도 '읽지 못했다'로 알린다
        raise SourceError(f"{Path(path).name} 의 내용이 형식에 맞지 않는다 ({type(exc).__name__}: {exc})") from None
    view.update(id=source_id(path), kind="light" if view.pop("light", False) else kind, path=str(path), name=Path(path).name)
    return view
